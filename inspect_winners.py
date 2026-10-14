import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

for filename in ["GANDHI_JAYANTI_LEADERBOARD_TSHIRT_WINNERS.xlsx", "TOP_3_WINNERS_TSHIRT_SIZES.xlsx"]:
    wb = openpyxl.load_workbook(filename)
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        print(f"\n=================== {filename} -> {sheet_name} ===================")
        for r in range(1, sheet.max_row + 1):
            row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
            if any(v is not None for v in row_vals):
                print(f"Row {r:2d}: {row_vals}")
