import openpyxl
import os
import sys

# Ensure utf-8 output
sys.stdout.reconfigure(encoding='utf-8')

files = [
    "GANDHI_JAYANTI_LEADERBOARD_TSHIRT_WINNERS.xlsx",
    "TOP_3_WINNERS_TSHIRT_SIZES.xlsx",
    "Complete gandhi jayanti.xlsx"
]

for filename in files:
    if not os.path.exists(filename):
        print(f"File not found: {filename}")
        continue
    print(f"\n=================== {filename} ===================")
    wb = openpyxl.load_workbook(filename)
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        print(f"\n--- Sheet: {sheet_name} (Rows: {sheet.max_row}, Cols: {sheet.max_column}) ---")
        for r in range(1, min(sheet.max_row + 1, 35)):
            row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
            if any(v is not None for v in row_vals):
                print(f"Row {r}: {row_vals}")
