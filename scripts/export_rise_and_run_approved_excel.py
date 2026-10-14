#!/usr/bin/env python3
"""
Mountain Run — Rise and Run Virtual Challenge 2026
Approved Proof Medal Dispatch List Generator

Generates a premium, courier-ready multi-sheet Excel workbook
for dispatching finisher medals to all 117 approved participants.
"""

import os
import sys
import datetime
from collections import Counter
import psycopg2
from psycopg2.extras import RealDictCursor
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

# Ensure UTF-8 output on Windows
if sys.platform.startswith("win"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

# Add root directory to path
sys.path.insert(0, os.path.abspath("."))
import scripts.medal_dispatch_manager as mm


def format_finish_time(seconds):
    if not seconds:
        return "N/A"
    try:
        sec = int(seconds)
        h = sec // 3600
        m = (sec % 3600) // 60
        s = sec % 60
        if h > 0:
            return f"{h:02d}:{m:02d}:{s:02d}"
        return f"{m:02d}:{s:02d}"
    except Exception:
        return str(seconds)


def clean_phone(phone_str):
    if not phone_str:
        return ""
    clean = str(phone_str).replace("+91", "").replace("-", "").replace(" ", "").strip()
    if clean.startswith("91") and len(clean) == 12:
        clean = clean[2:]
    return clean[-10:] if len(clean) >= 10 else clean


def standardize_distance(dist_str):
    if not dist_str:
        return "N/A"
    d = str(dist_str).strip().lower()
    if "21" in d:
        return "21 KM (Half Marathon)"
    elif "10" in d:
        return "10 KM"
    elif "5" in d:
        return "5 KM"
    elif "3.2" in d:
        return "3.2 KM"
    elif "1.6" in d:
        return "1.6 KM"
    return dist_str.upper()


def main():
    print("=" * 75)
    print("  MOUNTAIN RUN — RISE AND RUN MEDAL DISPATCH EXCEL GENERATOR")
    print("=" * 75)

    conn = mm.get_db_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)

    # 1. Fetch Rise and Run Event
    cur.execute("""
        SELECT id, title, slug, "startsAt", "endsAt"
        FROM "Event"
        WHERE slug = 'rise-and-run' OR title ILIKE '%Rise%Run%'
        LIMIT 1
    """)
    event = cur.fetchone()
    if not event:
        raise ValueError("Rise and Run event not found in database!")

    event_id = event["id"]
    event_title = event["title"]
    print(f"[*] Found Event: {event_title} (ID: {event_id})")

    # 2. Query all APPROVED registrations
    cur.execute("""
        SELECT 
            r.id as registration_id,
            r."bibNumber" as bib_number,
            r."userId" as user_id,
            u.name as user_name,
            u.email as email,
            u.phone as user_phone,
            r."shippingName" as shipping_name,
            r."shippingPhone" as shipping_phone,
            r."shippingLine1" as address_line1,
            r."shippingLine2" as address_line2,
            r."shippingCity" as city,
            r."shippingState" as state,
            r."shippingPincode" as pincode,
            r.distance as distance,
            r."activityType" as activity_type,
            r."finishTimeSeconds" as finish_time_seconds,
            r."proofStatus" as proof_status,
            r."registeredAt" as registered_at,
            COALESCE(m.status::text, 'PENDING') as medal_status,
            m."trackingNumber" as tracking_number,
            m.courier as courier,
            m."dispatchedAt" as dispatched_at,
            c."certificateNumber" as certificate_number,
            pu."activityImageUrl" as proof_url,
            pu."submittedAt" as proof_submitted_at,
            p."razorpayPaymentId" as payment_id,
            p."amountInPaise" as amount_in_paise,
            p."paidAt" as paid_at
        FROM "Registration" r
        JOIN "User" u ON r."userId" = u.id
        LEFT JOIN "Payment" p ON p."registrationId" = r.id
        LEFT JOIN "MedalDelivery" m ON m."registrationId" = r.id
        LEFT JOIN "ProofUpload" pu ON pu."registrationId" = r.id
        LEFT JOIN "Certificate" c ON c."registrationId" = r.id
        WHERE r."eventId" = %s
          AND r."proofStatus" = 'APPROVED'
        ORDER BY r."registeredAt" ASC
    """, (event_id,))

    rows = cur.fetchall()
    conn.close()

    print(f"[*] Total Approved Registrations Retrieved: {len(rows)}")

    # Deduplicate by BIB if necessary
    by_bib = {}
    for r in rows:
        bib = r["bib_number"]
        if bib not in by_bib:
            by_bib[bib] = r
        else:
            if not by_bib[bib]["proof_url"] and r["proof_url"]:
                by_bib[bib] = r

    athletes = list(by_bib.values())
    total_athletes = len(athletes)
    print(f"[*] Total Unique Athletes Ready for Dispatch: {total_athletes}")

    # Process and clean data
    processed_data = []
    for idx, r in enumerate(athletes, 1):
        line1 = str(r["address_line1"] or "").strip()
        line2 = str(r["address_line2"] or "").strip()
        city = str(r["city"] or "").strip().title()
        state = str(r["state"] or "").strip().title()
        pincode = str(r["pincode"] or "").strip()

        addr_parts = [p for p in [line1, line2] if p]
        full_street = ", ".join(addr_parts)
        full_shipping_addr = f"{full_street}, {city}, {state} - {pincode}".strip(" ,-")

        ship_phone = clean_phone(r["shipping_phone"])
        usr_phone = clean_phone(r["user_phone"])
        primary_phone = ship_phone or usr_phone
        alt_phone = usr_phone if usr_phone and usr_phone != ship_phone else ""

        recipient_name = (r["shipping_name"] or r["user_name"] or "Runner").strip().title()
        runner_name = (r["user_name"] or r["shipping_name"] or "Runner").strip().title()

        dist_clean = standardize_distance(r["distance"])
        activity = (r["activity_type"] or "running").strip().capitalize()
        finish_time = format_finish_time(r["finish_time_seconds"])

        reg_date = r["registered_at"].strftime("%d-%b-%Y") if r["registered_at"] else "N/A"
        proof_date = r["proof_submitted_at"].strftime("%d-%b-%Y %I:%M %p") if r["proof_submitted_at"] else "N/A"
        amount_rs = round((r["amount_in_paise"] or 0) / 100) if r["amount_in_paise"] else 0

        processed_data.append({
            "s_no": idx,
            "bib": r["bib_number"],
            "recipient_name": recipient_name,
            "runner_name": runner_name,
            "primary_phone": primary_phone,
            "alt_phone": alt_phone,
            "email": r["email"] or "",
            "full_address": full_shipping_addr,
            "address_line1": line1,
            "address_line2": line2,
            "city": city,
            "state": state,
            "pincode": pincode,
            "distance": dist_clean,
            "activity": activity,
            "finish_time": finish_time,
            "proof_status": r["proof_status"],
            "proof_date": proof_date,
            "proof_url": r["proof_url"] or "",
            "medal_status": (r["medal_status"] or "PENDING").upper(),
            "tracking_number": r["tracking_number"] or "",
            "courier": r["courier"] or "",
            "cert_number": r["certificate_number"] or "N/A",
            "reg_date": reg_date,
            "amount_rs": amount_rs,
            "payment_id": r["payment_id"] or "N/A"
        })

    # 3. Build the Excel Workbook with openpyxl
    wb = openpyxl.Workbook()

    # Define Theme Styles
    # Premium Mountain Run Forest Green & Sunrise Gold Palette
    COLOR_PRIMARY_DARK  = "0F3D24"    # Dark Forest Green (Title Banner)
    COLOR_PRIMARY_MED   = "1B633D"    # Medium Forest Green (Headers)
    COLOR_ZEBRA_ROW     = "F6FAF7"    # Very soft subtle mint zebra row
    COLOR_WHITE         = "FFFFFF"    # Crisp White
    COLOR_BORDER        = "D1E2D8"    # Soft Green-Gray Border
    COLOR_BORDER_DARK   = "0F3D24"    # Header Border
    COLOR_TEXT_MAIN     = "111827"    # Slate 900
    COLOR_TEXT_MUTED    = "4B5563"    # Slate 600

    # Pill fills and fonts
    FILL_APPROVED_BG    = "DEF7EC"    # Light Green
    FONT_APPROVED_FG    = "03543F"    # Dark Green
    FILL_PENDING_BG     = "FEF3C7"    # Light Amber
    FONT_PENDING_FG     = "92400E"    # Dark Amber
    FILL_PREPAID_BG     = "DBEAFE"    # Light Blue
    FONT_PREPAID_FG     = "1E40AF"    # Dark Blue

    thin_border = Border(
        left=Side(style="thin", color=COLOR_BORDER),
        right=Side(style="thin", color=COLOR_BORDER),
        top=Side(style="thin", color=COLOR_BORDER),
        bottom=Side(style="thin", color=COLOR_BORDER)
    )

    header_border = Border(
        left=Side(style="thin", color="0D4A2B"),
        right=Side(style="thin", color="0D4A2B"),
        top=Side(style="medium", color="0D4A2B"),
        bottom=Side(style="medium", color="0D4A2B")
    )

    f_banner_title = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    f_banner_sub = Font(name="Segoe UI", size=10, italic=True, color="E2E8F0")
    f_table_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    f_data_regular = Font(name="Segoe UI", size=9.5, color=COLOR_TEXT_MAIN)
    f_data_bold = Font(name="Segoe UI", size=9.5, bold=True, color=COLOR_TEXT_MAIN)
    f_data_muted = Font(name="Segoe UI", size=9, italic=True, color=COLOR_TEXT_MUTED)
    f_hyperlink = Font(name="Segoe UI", size=9.5, color="1D4ED8", underline="single")

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=False)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=False)
    align_wrap_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    fill_banner = PatternFill(start_color=COLOR_PRIMARY_DARK, end_color=COLOR_PRIMARY_DARK, fill_type="solid")
    fill_sub_banner = PatternFill(start_color=COLOR_PRIMARY_MED, end_color=COLOR_PRIMARY_MED, fill_type="solid")
    fill_header = PatternFill(start_color=COLOR_PRIMARY_MED, end_color=COLOR_PRIMARY_MED, fill_type="solid")
    fill_zebra = PatternFill(start_color=COLOR_ZEBRA_ROW, end_color=COLOR_ZEBRA_ROW, fill_type="solid")
    fill_white = PatternFill(start_color=COLOR_WHITE, end_color=COLOR_WHITE, fill_type="solid")
    fill_approved = PatternFill(start_color=FILL_APPROVED_BG, end_color=FILL_APPROVED_BG, fill_type="solid")
    fill_pending = PatternFill(start_color=FILL_PENDING_BG, end_color=FILL_PENDING_BG, fill_type="solid")
    fill_prepaid = PatternFill(start_color=FILL_PREPAID_BG, end_color=FILL_PREPAID_BG, fill_type="solid")

    # =========================================================================
    # SHEET 1: COURIER DISPATCH READY (Customized for 3PL / Logistics Delivery)
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Courier Dispatch Ready"
    ws1.views.sheetView[0].showGridLines = True
    ws1.freeze_panes = "D4"

    # Header Definition for Sheet 1
    COURIER_COLS = [
        ("S.No", 6, align_center),
        ("BIB Number", 13, align_center),
        ("Recipient Name", 22, align_left),
        ("Primary Mobile (10 Digits)", 16, align_center),
        ("Alternate Mobile", 16, align_center),
        ("Email Address", 26, align_left),
        ("Address Line 1", 30, align_left),
        ("Address Line 2 / Landmark", 22, align_left),
        ("Full Shipping Address", 44, align_wrap_left),
        ("City", 16, align_left),
        ("State", 18, align_left),
        ("Pincode", 10, align_center),
        ("Package Contents", 28, align_left),
        ("Category", 18, align_center),
        ("Package Weight", 14, align_center),
        ("Payment Mode", 14, align_center),
        ("Medal Status", 16, align_center),
        ("Courier Partner", 16, align_center),
        ("AWB / Tracking No", 20, align_center),
        ("Dispatch Date", 14, align_center),
        ("Delivery Status", 14, align_center),
        ("Remarks / Notes", 20, align_left),
    ]

    last_col_letter_ws1 = get_column_letter(len(COURIER_COLS))

    # Row 1: Title Banner
    ws1.merge_cells(f"A1:{last_col_letter_ws1}1")
    b1 = ws1["A1"]
    b1.value = "MOUNTAIN RUN  |  RISE AND RUN VIRTUAL CHALLENGE 2026"
    b1.font = f_banner_title
    b1.fill = fill_banner
    b1.alignment = align_center
    ws1.row_dimensions[1].height = 32

    # Row 2: Subtitle Banner
    now_str = datetime.datetime.now().strftime("%d-%b-%Y %I:%M %p")
    ws1.merge_cells(f"A2:{last_col_letter_ws1}2")
    b2 = ws1["A2"]
    b2.value = f"Official Courier Dispatch Sheet  |  {total_athletes} Approved Athletes Ready for Medal Delivery  |  Generated: {now_str}"
    b2.font = f_banner_sub
    b2.fill = fill_sub_banner
    b2.alignment = align_center
    ws1.row_dimensions[2].height = 20

    # Row 3: Table Headers
    ws1.row_dimensions[3].height = 26
    for c_idx, (col_name, col_width, col_align) in enumerate(COURIER_COLS, 1):
        c_letter = get_column_letter(c_idx)
        ws1.column_dimensions[c_letter].width = col_width
        cell = ws1.cell(row=3, column=c_idx, value=col_name)
        cell.font = f_table_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = header_border

    # Data Rows
    for row_idx, item in enumerate(processed_data, start=4):
        is_even = (row_idx % 2 == 0)
        row_fill = fill_zebra if is_even else fill_white
        ws1.row_dimensions[row_idx].height = 22

        row_values = [
            (item["s_no"], align_center, f_data_regular),
            (item["bib"], align_center, f_data_bold),
            (item["recipient_name"], align_left, f_data_bold),
            (item["primary_phone"], align_center, f_data_regular),
            (item["alt_phone"], align_center, f_data_regular),
            (item["email"], align_left, f_data_regular),
            (item["address_line1"], align_left, f_data_regular),
            (item["address_line2"], align_left, f_data_regular),
            (item["full_address"], align_wrap_left, f_data_regular),
            (item["city"], align_left, f_data_regular),
            (item["state"], align_left, f_data_regular),
            (item["pincode"], align_center, f_data_bold),
            ("Rise and Run Official Finisher Medal", align_left, f_data_regular),
            (item["distance"], align_center, f_data_regular),
            ("0.15 kg", align_center, f_data_regular),
            ("PREPAID", align_center, Font(name="Segoe UI", size=9, bold=True, color=FONT_PREPAID_FG)),
            ("PENDING DISPATCH", align_center, Font(name="Segoe UI", size=9, bold=True, color=FONT_PENDING_FG)),
            (item["courier"], align_center, f_data_regular),
            (item["tracking_number"], align_center, f_data_regular),
            ("", align_center, f_data_regular), # Dispatch Date
            ("", align_center, f_data_regular), # Delivery Status
            ("", align_left, f_data_regular),   # Remarks
        ]

        for c_idx, (val, cell_align, cell_font) in enumerate(row_values, 1):
            cell = ws1.cell(row=row_idx, column=c_idx, value=val if val != "" else None)
            cell.font = cell_font
            cell.alignment = cell_align
            cell.border = thin_border
            cell.fill = row_fill

            # Format Pincode & Phones explicitly as text
            if c_idx in [4, 5, 12]:
                cell.number_format = "@"

            # Payment Mode Pill
            if c_idx == 16:
                cell.fill = fill_prepaid

            # Medal Status Pill
            if c_idx == 17:
                cell.fill = fill_pending

    ws1.auto_filter.ref = f"A3:{last_col_letter_ws1}{3 + total_athletes}"

    # Print setup
    ws1.print_title_rows = "1:3"
    ws1.page_setup.orientation = ws1.ORIENTATION_LANDSCAPE
    ws1.page_setup.paperSize = ws1.PAPERSIZE_A4
    ws1.page_setup.fitToPage = True
    ws1.page_setup.fitToWidth = 1
    ws1.page_setup.fitToHeight = 0

    # =========================================================================
    # SHEET 2: MASTER DISPATCH & PROOFS (Full Operational Records)
    # =========================================================================
    ws2 = wb.create_sheet("Master Data & Proofs")
    ws2.views.sheetView[0].showGridLines = True
    ws2.freeze_panes = "D4"

    MASTER_COLS = [
        ("S.No", 6, align_center),
        ("BIB Number", 13, align_center),
        ("Runner Name", 22, align_left),
        ("Shipping Recipient", 22, align_left),
        ("Mobile Number", 14, align_center),
        ("Email Address", 26, align_left),
        ("Distance", 18, align_center),
        ("Activity", 11, align_center),
        ("Finish Time", 12, align_center),
        ("Full Shipping Address", 44, align_wrap_left),
        ("Address Line 1", 28, align_left),
        ("Address Line 2", 20, align_left),
        ("City", 16, align_left),
        ("State", 18, align_left),
        ("Pincode", 10, align_center),
        ("Proof Status", 14, align_center),
        ("Proof Submitted Date", 20, align_center),
        ("Proof Screenshot", 18, align_center),
        ("Medal Status", 16, align_center),
        ("Courier Partner", 15, align_center),
        ("AWB / Tracking No", 18, align_center),
        ("Certificate Number", 22, align_center),
        ("Registration Date", 14, align_center),
        ("Amount (Rs)", 12, align_right),
        ("Payment ID", 22, align_center),
    ]

    last_col_letter_ws2 = get_column_letter(len(MASTER_COLS))

    # Banner Row 1 & 2
    ws2.merge_cells(f"A1:{last_col_letter_ws2}1")
    mb1 = ws2["A1"]
    mb1.value = "MOUNTAIN RUN  |  RISE AND RUN: COMPLETE APPROVED PROOF & DISPATCH MASTER"
    mb1.font = f_banner_title
    mb1.fill = fill_banner
    mb1.alignment = align_center
    ws2.row_dimensions[1].height = 32

    ws2.merge_cells(f"A2:{last_col_letter_ws2}2")
    mb2 = ws2["A2"]
    mb2.value = f"Complete Audit Master  |  {total_athletes} Approved Records  |  Verified Activity Proof Screenshots  |  {now_str}"
    mb2.font = f_banner_sub
    mb2.fill = fill_sub_banner
    mb2.alignment = align_center
    ws2.row_dimensions[2].height = 20

    # Headers
    ws2.row_dimensions[3].height = 26
    for c_idx, (col_name, col_width, col_align) in enumerate(MASTER_COLS, 1):
        c_letter = get_column_letter(c_idx)
        ws2.column_dimensions[c_letter].width = col_width
        cell = ws2.cell(row=3, column=c_idx, value=col_name)
        cell.font = f_table_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = header_border

    # Data Rows
    for row_idx, item in enumerate(processed_data, start=4):
        is_even = (row_idx % 2 == 0)
        row_fill = fill_zebra if is_even else fill_white
        ws2.row_dimensions[row_idx].height = 22

        m_values = [
            (item["s_no"], align_center, f_data_regular),
            (item["bib"], align_center, f_data_bold),
            (item["runner_name"], align_left, f_data_bold),
            (item["recipient_name"], align_left, f_data_regular),
            (item["primary_phone"], align_center, f_data_regular),
            (item["email"], align_left, f_data_regular),
            (item["distance"], align_center, f_data_regular),
            (item["activity"], align_center, f_data_regular),
            (item["finish_time"], align_center, f_data_regular),
            (item["full_address"], align_wrap_left, f_data_regular),
            (item["address_line1"], align_left, f_data_regular),
            (item["address_line2"], align_left, f_data_regular),
            (item["city"], align_left, f_data_regular),
            (item["state"], align_left, f_data_regular),
            (item["pincode"], align_center, f_data_bold),
            ("APPROVED", align_center, Font(name="Segoe UI", size=9, bold=True, color=FONT_APPROVED_FG)),
            (item["proof_date"], align_center, f_data_regular),
            ("View Proof", align_center, f_hyperlink), # Clickable URL
            ("PENDING", align_center, Font(name="Segoe UI", size=9, bold=True, color=FONT_PENDING_FG)),
            (item["courier"], align_center, f_data_regular),
            (item["tracking_number"], align_center, f_data_regular),
            (item["cert_number"], align_center, f_data_regular),
            (item["reg_date"], align_center, f_data_regular),
            (item["amount_rs"], align_right, f_data_regular),
            (item["payment_id"], align_center, f_data_regular),
        ]

        for c_idx, (val, cell_align, cell_font) in enumerate(m_values, 1):
            cell = ws2.cell(row=row_idx, column=c_idx, value=val if val != "" else None)
            cell.font = cell_font
            cell.alignment = cell_align
            cell.border = thin_border
            cell.fill = row_fill

            # Format Pincode & Phone
            if c_idx in [5, 15]:
                cell.number_format = "@"
            elif c_idx == 24:
                cell.number_format = "₹#,##0"

            # Clickable Proof link
            if c_idx == 18 and item["proof_url"]:
                cell.hyperlink = item["proof_url"]

            # Proof Status Pill
            if c_idx == 16:
                cell.fill = fill_approved

            # Medal Status Pill
            if c_idx == 19:
                cell.fill = fill_pending

    ws2.auto_filter.ref = f"A3:{last_col_letter_ws2}{3 + total_athletes}"
    ws2.print_title_rows = "1:3"
    ws2.page_setup.orientation = ws2.ORIENTATION_LANDSCAPE
    ws2.page_setup.paperSize = ws2.PAPERSIZE_A4

    # =========================================================================
    # SHEET 3: SUMMARY & LOGISTICS DASHBOARD
    # =========================================================================
    ws3 = wb.create_sheet("Summary & Analytics")
    ws3.views.sheetView[0].showGridLines = True

    ws3.column_dimensions["A"].width = 5
    ws3.column_dimensions["B"].width = 28
    ws3.column_dimensions["C"].width = 16
    ws3.column_dimensions["D"].width = 16
    ws3.column_dimensions["E"].width = 5
    ws3.column_dimensions["F"].width = 28
    ws3.column_dimensions["G"].width = 16
    ws3.column_dimensions["H"].width = 16

    # Banner
    ws3.merge_cells("B1:H1")
    db1 = ws3["B1"]
    db1.value = "MOUNTAIN RUN  |  RISE AND RUN DISPATCH DASHBOARD"
    db1.font = f_banner_title
    db1.fill = fill_banner
    db1.alignment = align_center
    ws3.row_dimensions[1].height = 32

    ws3.merge_cells("B2:H2")
    db2 = ws3["B2"]
    db2.value = f"Operational Metrics, Category Breakdown & Logistics Planning  |  Generated {now_str}"
    db2.font = f_banner_sub
    db2.fill = fill_sub_banner
    db2.alignment = align_center
    ws3.row_dimensions[2].height = 20

    # Key Metrics Cards (Row 4 - 5)
    # Card 1: Total Finishers (B4:C5)
    ws3.merge_cells("B4:C4")
    ws3["B4"].value = "TOTAL APPROVED FINISHERS"
    ws3["B4"].font = Font(name="Segoe UI", size=9, bold=True, color="065F46")
    ws3["B4"].fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    ws3["B4"].alignment = align_center

    ws3.merge_cells("B5:C5")
    ws3["B5"].value = total_athletes
    ws3["B5"].font = Font(name="Segoe UI", size=20, bold=True, color="065F46")
    ws3["B5"].fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    ws3["B5"].alignment = align_center
    ws3["B4"].border = thin_border
    ws3["B5"].border = thin_border
    ws3["C4"].border = thin_border
    ws3["C5"].border = thin_border

    # Card 2: Pending Dispatch (D4:E4, D5:E5)
    ws3.merge_cells("D4:E4")
    ws3["D4"].value = "MEDALS PENDING DISPATCH"
    ws3["D4"].font = Font(name="Segoe UI", size=9, bold=True, color="92400E")
    ws3["D4"].fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    ws3["D4"].alignment = align_center

    ws3.merge_cells("D5:E5")
    ws3["D5"].value = total_athletes
    ws3["D5"].font = Font(name="Segoe UI", size=20, bold=True, color="92400E")
    ws3["D5"].fill = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")
    ws3["D5"].alignment = align_center
    ws3["D4"].border = thin_border
    ws3["D5"].border = thin_border
    ws3["E4"].border = thin_border
    ws3["E5"].border = thin_border

    # Card 3: Indian States Covered (F4:H4, F5:H5)
    state_count = len(set(r["state"] for r in processed_data if r["state"]))
    city_count = len(set(r["city"] for r in processed_data if r["city"]))
    ws3.merge_cells("F4:H4")
    ws3["F4"].value = "DESTINATION REACH"
    ws3["F4"].font = Font(name="Segoe UI", size=9, bold=True, color="1E40AF")
    ws3["F4"].fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    ws3["F4"].alignment = align_center

    ws3.merge_cells("F5:H5")
    ws3["F5"].value = f"{state_count} States  |  {city_count} Cities Across India"
    ws3["F5"].font = Font(name="Segoe UI", size=15, bold=True, color="1E40AF")
    ws3["F5"].fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    ws3["F5"].alignment = align_center
    for c in ["F", "G", "H"]:
        ws3[f"{c}4"].border = thin_border
        ws3[f"{c}5"].border = thin_border

    ws3.row_dimensions[4].height = 18
    ws3.row_dimensions[5].height = 30

    # Section Headers helper
    def add_section_header(row, col_start, col_end, title):
        ws3.merge_cells(f"{col_start}{row}:{col_end}{row}")
        c = ws3[f"{col_start}{row}"]
        c.value = title
        c.font = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color=COLOR_PRIMARY_MED, end_color=COLOR_PRIMARY_MED, fill_type="solid")
        c.alignment = align_center
        ws3.row_dimensions[row].height = 22

    # Left Column: Distance Category Breakdown (Row 7)
    add_section_header(7, "B", "D", "BREAKDOWN BY DISTANCE CATEGORY")
    ws3.cell(row=8, column=2, value="Distance Category").font = f_table_header
    ws3.cell(row=8, column=2).fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    ws3.cell(row=8, column=2).alignment = align_left
    ws3.cell(row=8, column=2).border = header_border

    ws3.cell(row=8, column=3, value="Athletes").font = f_table_header
    ws3.cell(row=8, column=3).fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    ws3.cell(row=8, column=3).alignment = align_center
    ws3.cell(row=8, column=3).border = header_border

    ws3.cell(row=8, column=4, value="Share (%)").font = f_table_header
    ws3.cell(row=8, column=4).fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    ws3.cell(row=8, column=4).alignment = align_center
    ws3.cell(row=8, column=4).border = header_border
    ws3.row_dimensions[8].height = 20

    dist_counts = Counter(r["distance"] for r in processed_data)
    dist_order = ["5 KM", "10 KM", "1.6 KM", "21 KM (Half Marathon)", "3.2 KM"]
    d_row = 9
    for d_name in dist_order:
        cnt = dist_counts.get(d_name, 0)
        pct = (cnt / total_athletes) * 100 if total_athletes else 0
        is_e = (d_row % 2 == 0)
        rf = fill_zebra if is_e else fill_white

        c1 = ws3.cell(row=d_row, column=2, value=d_name)
        c2 = ws3.cell(row=d_row, column=3, value=cnt)
        c3 = ws3.cell(row=d_row, column=4, value=f"{pct:.1f}%")

        c1.font = f_data_bold
        c2.font = f_data_bold
        c3.font = f_data_regular
        c1.alignment = align_left
        c2.alignment = align_center
        c3.alignment = align_center
        c1.border = thin_border
        c2.border = thin_border
        c3.border = thin_border
        c1.fill = rf
        c2.fill = rf
        c3.fill = rf
        ws3.row_dimensions[d_row].height = 20
        d_row += 1

    # Total Row for Distances
    c1 = ws3.cell(row=d_row, column=2, value="Total Medals")
    c2 = ws3.cell(row=d_row, column=3, value=total_athletes)
    c3 = ws3.cell(row=d_row, column=4, value="100.0%")
    c1.font = Font(name="Segoe UI", size=10, bold=True, color="065F46")
    c2.font = Font(name="Segoe UI", size=10, bold=True, color="065F46")
    c3.font = Font(name="Segoe UI", size=10, bold=True, color="065F46")
    c1.alignment = align_left
    c2.alignment = align_center
    c3.alignment = align_center
    c1.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    c2.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    c3.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    c1.border = header_border
    c2.border = header_border
    c3.border = header_border
    ws3.row_dimensions[d_row].height = 22

    # Right Column: Top Delivery States (Row 7)
    add_section_header(7, "F", "H", "TOP DELIVERY STATES DISTRIBUTION")
    ws3.cell(row=8, column=6, value="State / Union Territory").font = f_table_header
    ws3.cell(row=8, column=6).fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    ws3.cell(row=8, column=6).alignment = align_left
    ws3.cell(row=8, column=6).border = header_border

    ws3.cell(row=8, column=7, value="Medals").font = f_table_header
    ws3.cell(row=8, column=7).fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    ws3.cell(row=8, column=7).alignment = align_center
    ws3.cell(row=8, column=7).border = header_border

    ws3.cell(row=8, column=8, value="Share (%)").font = f_table_header
    ws3.cell(row=8, column=8).fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    ws3.cell(row=8, column=8).alignment = align_center
    ws3.cell(row=8, column=8).border = header_border

    state_counts = Counter(r["state"] for r in processed_data if r["state"])
    s_row = 9
    for state_name, cnt in state_counts.most_common(12):
        pct = (cnt / total_athletes) * 100 if total_athletes else 0
        is_e = (s_row % 2 == 0)
        rf = fill_zebra if is_e else fill_white

        c1 = ws3.cell(row=s_row, column=6, value=state_name)
        c2 = ws3.cell(row=s_row, column=7, value=cnt)
        c3 = ws3.cell(row=s_row, column=8, value=f"{pct:.1f}%")

        c1.font = f_data_bold
        c2.font = f_data_bold
        c3.font = f_data_regular
        c1.alignment = align_left
        c2.alignment = align_center
        c3.alignment = align_center
        c1.border = thin_border
        c2.border = thin_border
        c3.border = thin_border
        c1.fill = rf
        c2.fill = rf
        c3.fill = rf
        ws3.row_dimensions[s_row].height = 20
        s_row += 1

    # Remaining states summary
    top_sum = sum(cnt for _, cnt in state_counts.most_common(12))
    other_sum = total_athletes - top_sum
    if other_sum > 0:
        c1 = ws3.cell(row=s_row, column=6, value=f"Other States ({len(state_counts) - 12} states)")
        c2 = ws3.cell(row=s_row, column=7, value=other_sum)
        c3 = ws3.cell(row=s_row, column=8, value=f"{(other_sum / total_athletes)*100:.1f}%")
        c1.font = f_data_muted
        c2.font = f_data_bold
        c3.font = f_data_regular
        c1.alignment = align_left
        c2.alignment = align_center
        c3.alignment = align_center
        c1.border = thin_border
        c2.border = thin_border
        c3.border = thin_border
        c1.fill = fill_white
        c2.fill = fill_white
        c3.fill = fill_white
        ws3.row_dimensions[s_row].height = 20
        s_row += 1

    # Logistics & Packaging Instructions (Below distance table)
    instr_row = d_row + 2
    add_section_header(instr_row, "B", "D", "LOGISTICS & PACKAGING PROTOCOL")
    instr_row += 1

    instructions = [
        "1. Medal Packaging: Verify 'Rise and Run' medal is undamaged.",
        "2. Protection: Wrap in bubble pouch & seal in tear-resistant polybag.",
        "3. Address Label: Stick standard shipping label with Barcode / AWB.",
        "4. Contact Number: Ensure 10-digit mobile is clearly printed for OTP.",
        "5. Courier Partners: India Post Speed Post / Delhivery / DTDC.",
        "6. Tracking Update: Enter courier & AWB back into sheet once dispatched."
    ]

    for line in instructions:
        ws3.merge_cells(f"B{instr_row}:D{instr_row}")
        c = ws3[f"B{instr_row}"]
        c.value = line
        c.font = Font(name="Segoe UI", size=9, color="374151")
        c.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
        c.alignment = align_left
        c.border = thin_border
        ws3.cell(row=instr_row, column=3).border = thin_border
        ws3.cell(row=instr_row, column=4).border = thin_border
        ws3.row_dimensions[instr_row].height = 19
        instr_row += 1

    # Save Workbook
    output_filename = "RISE_AND_RUN_MEDAL_DISPATCH_LIST.xlsx"
    output_path = os.path.join(os.getcwd(), output_filename)
    wb.save(output_path)

    print("\n" + "=" * 75)
    print("SUCCESS: EXCEL SHEET GENERATED!")
    print(f"File Path: {output_path}")
    print("Sheets:")
    print(f"  1. 'Courier Dispatch Ready'  -> {total_athletes} rows, clean 22 cols for 3PL/Courier")
    print(f"  2. 'Master Data & Proofs'    -> Full 25 cols with clickable proof links & cert info")
    print(f"  3. 'Summary & Analytics'     -> Key KPI cards, distance breakdown, state distribution")
    print("=" * 75)

    return output_path, total_athletes


if __name__ == "__main__":
    main()
