import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load the 21 podium winners
with open('scripts/top3_sports_day.json', 'r', encoding='utf-8') as f:
    winners = json.load(f)

# Load existing dispatch data if needed for all finishers
all_finishers = []
if os.path.exists('SPORTS_DAY_APPROVED_PROOFS_DISPATCH_LIST.xlsx'):
    try:
        old_wb = openpyxl.load_workbook('SPORTS_DAY_APPROVED_PROOFS_DISPATCH_LIST.xlsx')
        old_ws = old_wb.active
        headers = [c.value for c in old_ws[1]]
        for row in old_ws.iter_rows(min_row=2, values_only=True):
            if any(row):
                all_finishers.append(dict(zip(headers, row)))
    except Exception as e:
        print("Warning reading old dispatch list:", e)

# -------------------------------------------------------------
# CREATE THE SINGLE MASTER LEADERBOARD SHEET
# -------------------------------------------------------------
wb = openpyxl.Workbook()

# Styles
font_title = Font(name="Calibri", size=15, bold=True, color="1B4D3E")
font_sub = Font(name="Calibri", size=10, italic=True, color="555555")
font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid")
fill_sub_header = PatternFill(start_color="2D6A4F", end_color="2D6A4F", fill_type="solid")

font_gold = Font(name="Calibri", size=10, bold=True, color="92400E")
fill_gold = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

font_silver = Font(name="Calibri", size=10, bold=True, color="374151")
fill_silver = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

font_bronze = Font(name="Calibri", size=10, bold=True, color="7C2D12")
fill_bronze = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")

font_data = Font(name="Calibri", size=10)
font_bold_data = Font(name="Calibri", size=10, bold=True)
align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")

thin_border = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB')
)

# -------------------------------------------------------------
# SHEET 1: Top 3 T-Shirt Winners (21 Athletes)
# -------------------------------------------------------------
ws1 = wb.active
ws1.title = "Top 3 T-Shirt Winners"

ws1.merge_cells("A1:N1")
ws1["A1"] = "MOUNTAIN RUN — SPORTS DAY CELEBRATION 2026"
ws1["A1"].font = font_title
ws1["A1"].alignment = align_center

ws1.merge_cells("A2:N2")
ws1["A2"] = "Official Category Podium Finishers (Rank 1, 2, 3) — Free T-Shirt Dispatch List (21 Athletes)"
ws1["A2"].font = font_sub
ws1["A2"].alignment = align_center

headers1 = [
    "S.No",
    "Category",
    "Rank",
    "BIB Number",
    "Runner Name",
    "Finish Time",
    "Mobile Number",
    "Email Address",
    "Full Shipping Address",
    "City",
    "State",
    "Pincode",
    "T-Shirt Size",
    "Dispatch Status"
]

ws1.row_dimensions[1].height = 28
ws1.row_dimensions[2].height = 18
ws1.row_dimensions[4].height = 24

for col_idx, h in enumerate(headers1, 1):
    cell = ws1.cell(row=4, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = thin_border

for idx, w in enumerate(winners, 1):
    row_num = 4 + idx
    ws1.row_dimensions[row_num].height = 22
    
    rank_label = f"1st Place" if w['rank'] == 1 else (f"2nd Place" if w['rank'] == 2 else f"3rd Place")
    
    full_loc = f"{w['city']}, {w['state']} - {w['pincode']}" if w['city'] else ""
    full_addr = f"{w['address']}, {full_loc}".strip().rstrip(',')
    
    row_vals = [
        idx,
        w['category'],
        rank_label,
        w['bib'],
        w['name'],
        w['time'],
        w['phone'],
        w['email'],
        full_addr,
        w['city'],
        w['state'],
        w['pincode'],
        "", # Blank for size entry when user replies
        "Pending Size Confirmation"
    ]
    
    for c_idx, val in enumerate(row_vals, 1):
        c = ws1.cell(row=row_num, column=c_idx, value=val)
        c.font = font_data
        c.border = thin_border
        
        if c_idx in [1, 2, 3, 4, 6, 12, 13, 14]:
            c.alignment = align_center
        else:
            c.alignment = align_left
            
        if c_idx == 3:
            if w['rank'] == 1:
                c.fill = fill_gold
                c.font = font_gold
            elif w['rank'] == 2:
                c.fill = fill_silver
                c.font = font_silver
            elif w['rank'] == 3:
                c.fill = fill_bronze
                c.font = font_bronze

for col in ws1.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.row < 4:
            continue
        max_len = max(max_len, len(str(cell.value or "")))
    ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)
ws1.column_dimensions['I'].width = 45
ws1.column_dimensions['M'].width = 16
ws1.column_dimensions['N'].width = 24

# -------------------------------------------------------------
# SHEET 2: Category Summary Dashboard
# -------------------------------------------------------------
ws2 = wb.create_sheet(title="Category Summary")

ws2.merge_cells("A1:F1")
ws2["A1"] = "SPORTS DAY CELEBRATION — CATEGORY PODIUM SUMMARY"
ws2["A1"].font = font_title
ws2["A1"].alignment = align_center

ws2.row_dimensions[1].height = 28
ws2.row_dimensions[3].height = 24

headers2 = ["Category", "1st Place Winner", "2nd Place Winner", "3rd Place Winner", "Total Winners", "T-Shirts Allocated"]
for col_idx, h in enumerate(headers2, 1):
    cell = ws2.cell(row=3, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = thin_border

# Group winners by category
cats = ["1.6 km", "3.2 km", "5 km", "7 km", "10 km", "15 km", "21 km"]
cat_winners = {}
for w in winners:
    cat_winners.setdefault(w['category'], []).append(w)

for idx, cat in enumerate(cats, 1):
    row_num = 3 + idx
    ws2.row_dimensions[row_num].height = 22
    w_list = cat_winners.get(cat, [])
    
    r1 = f"{w_list[0]['name']} ({w_list[0]['time']})" if len(w_list) >= 1 else "N/A"
    r2 = f"{w_list[1]['name']} ({w_list[1]['time']})" if len(w_list) >= 2 else "N/A"
    r3 = f"{w_list[2]['name']} ({w_list[2]['time']})" if len(w_list) >= 3 else "N/A"
    
    row_vals = [cat, r1, r2, r3, len(w_list), "3 T-Shirts"]
    for c_idx, val in enumerate(row_vals, 1):
        c = ws2.cell(row=row_num, column=c_idx, value=val)
        c.font = font_bold_data if c_idx in [1, 5, 6] else font_data
        c.border = thin_border
        c.alignment = align_center if c_idx in [1, 5, 6] else align_left

# Total Row
tot_row = 3 + len(cats) + 1
ws2.row_dimensions[tot_row].height = 24
tot_vals = ["TOTAL", "-", "-", "-", "21 Winners", "21 T-Shirts"]
for c_idx, val in enumerate(tot_vals, 1):
    c = ws2.cell(row=tot_row, column=c_idx, value=val)
    c.font = font_bold_data
    c.fill = fill_silver
    c.border = thin_border
    c.alignment = align_center

for col in ws2.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.row < 3:
            continue
        max_len = max(max_len, len(str(cell.value or "")))
    ws2.column_dimensions[col_letter].width = max(max_len + 4, 15)

# -------------------------------------------------------------
# SHEET 3: All Verified Athletes (Complete Registry)
# -------------------------------------------------------------
if all_finishers:
    ws3 = wb.create_sheet(title="All Athletes Archive")
    ws3.merge_cells("A1:K1")
    ws3["A1"] = "SPORTS DAY CELEBRATION — ALL VERIFIED PARTICIPANTS"
    ws3["A1"].font = font_title
    ws3["A1"].alignment = align_center
    
    headers3 = ["S.No", "BIB Number", "Runner Name", "Category", "Finish Time", "Mobile", "Email", "City", "State", "Pincode", "T-Shirt Eligible?"]
    ws3.row_dimensions[1].height = 28
    ws3.row_dimensions[3].height = 24
    
    for col_idx, h in enumerate(headers3, 1):
        cell = ws3.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_sub_header
        cell.alignment = align_center
        cell.border = thin_border
        
    top3_bibs = set(w['bib'] for w in winners)
    
    for idx, f_row in enumerate(all_finishers, 1):
        row_num = 3 + idx
        ws3.row_dimensions[row_num].height = 19
        bib = str(f_row.get('BIB Number') or '')
        is_tshirt = "YES (Top 3 Winner)" if bib in top3_bibs else "Medal & Certificate Only"
        
        row_vals = [
            idx,
            bib,
            f_row.get('Runner Name') or f_row.get('Shipping Recipient') or '',
            f_row.get('Category / Distance') or '',
            f_row.get('Finish Time') or '',
            f_row.get('Mobile Number') or '',
            f_row.get('Email Address') or '',
            f_row.get('City') or '',
            f_row.get('State') or '',
            f_row.get('Pincode') or '',
            is_tshirt
        ]
        for c_idx, val in enumerate(row_vals, 1):
            c = ws3.cell(row=row_num, column=c_idx, value=val)
            c.font = font_data
            c.border = thin_border
            if c_idx in [1, 2, 4, 5, 10, 11]:
                c.alignment = align_center
            else:
                c.alignment = align_left
            if is_tshirt.startswith("YES") and c_idx == 11:
                c.fill = fill_gold
                c.font = font_gold

    for col in ws3.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 3:
                continue
            max_len = max(max_len, len(str(cell.value or "")))
        ws3.column_dimensions[col_letter].width = max(max_len + 3, 12)

# Save Master File
output_master = "SPORTS_DAY_LEADERBOARD_TSHIRT_WINNERS.xlsx"
wb.save(output_master)
print(f"SUCCESS: Created single master excel sheet: {output_master}")

# Clean up old Excel sheets ONLY (no code or scripts)
old_sheets = [
    'PROOF_NOT_SUBMITTED.xlsx',
    'SPORTS_DAY_APPROVED_PROOFS_DISPATCH_LIST.xlsx',
    'SPORTS_DAY_DISPATCH_LIST_v2.xlsx',
    'SPORTS_DAY_TOP3_WINNERS_TSHIRT_LIST.xlsx'
]

deleted = []
for f in old_sheets:
    if os.path.exists(f) and f != output_master:
        try:
            os.remove(f)
            deleted.append(f)
            print(f"Deleted old sheet: {f}")
        except Exception as e:
            print(f"Failed to delete {f}: {e}")

print(f"Total deleted sheets: {len(deleted)}")
