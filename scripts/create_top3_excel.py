import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open('scripts/top3_sports_day.json', 'r', encoding='utf-8') as f:
    winners = json.load(f)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Top 3 T-Shirt Winners"

# Styles
font_title = Font(name="Calibri", size=16, bold=True, color="1B4D3E")
font_sub = Font(name="Calibri", size=11, italic=True, color="555555")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid")

font_gold = Font(name="Calibri", size=11, bold=True, color="92400E")
fill_gold = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

font_silver = Font(name="Calibri", size=11, bold=True, color="374151")
fill_silver = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

font_bronze = Font(name="Calibri", size=11, bold=True, color="7C2D12")
fill_bronze = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")

font_data = Font(name="Calibri", size=11)
align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")

thin_border = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB')
)

# Title Block
ws.merge_cells("A1:K1")
ws["A1"] = "MOUNTAIN RUN — SPORTS DAY CELEBRATION 2026"
ws["A1"].font = font_title
ws["A1"].alignment = align_center

ws.merge_cells("A2:K2")
ws["A2"] = "Official Top 3 Category Winners — T-Shirt Dispatch List (21 Athletes)"
ws["A2"].font = font_sub
ws["A2"].alignment = align_center

ws.row_dimensions[1].height = 30
ws.row_dimensions[2].height = 20
ws.row_dimensions[4].height = 26

# Headers
headers = [
    "S.No",
    "Category",
    "Rank",
    "BIB Number",
    "Runner Name",
    "Finish Time",
    "Mobile Number",
    "Email Address",
    "Full Shipping Address",
    "City & State",
    "T-Shirt Size (To Confirm)"
]

ws.append([]) # Row 3 blank
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = thin_border

# Populate Data
for idx, w in enumerate(winners, 1):
    row_num = 4 + idx
    ws.row_dimensions[row_num].height = 22
    
    rank_label = f"🥇 1st" if w['rank'] == 1 else (f"🥈 2nd" if w['rank'] == 2 else f"🥉 3rd")
    
    full_loc = f"{w['city']}, {w['state']} - {w['pincode']}" if w['city'] else ""
    full_addr = f"{w['address']}, {full_loc}".strip().rstrip(',')
    
    row_data = [
        idx,
        w['category'],
        rank_label,
        w['bib'],
        w['name'],
        w['time'],
        w['phone'],
        w['email'],
        full_addr,
        f"{w['city']}, {w['state']}",
        "L (Suggested)" # Placeholder / default
    ]
    
    for col_idx, val in enumerate(row_data, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=val)
        cell.font = font_data
        cell.border = thin_border
        
        if col_idx in [1, 2, 3, 4, 6, 11]:
            cell.alignment = align_center
        else:
            cell.alignment = align_left
            
        # Highlight rank
        if col_idx == 3:
            if w['rank'] == 1:
                cell.font = font_gold
                cell.fill = fill_gold
            elif w['rank'] == 2:
                cell.font = font_silver
                cell.fill = fill_silver
            elif w['rank'] == 3:
                cell.font = font_bronze
                cell.fill = fill_bronze

# Column auto width
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.row < 4:
            continue
        val_str = str(cell.value or "")
        max_len = max(max_len, len(val_str))
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

ws.column_dimensions['I'].width = 45 # Shipping address

output_path = "SPORTS_DAY_TOP3_WINNERS_TSHIRT_LIST.xlsx"
wb.save(output_path)
print(f"SUCCESS: Created Excel sheet: {output_path}")
