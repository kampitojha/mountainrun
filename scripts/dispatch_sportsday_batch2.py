import os
import sys
import json
import time
import datetime
import urllib.request
import urllib.error
import openpyxl
import psycopg2
from psycopg2.extras import RealDictCursor

sys.path.insert(0, os.path.abspath("."))
import scripts.medal_dispatch_manager as mm
from scripts.dispatch_sportsday_tracking_emails import build_sportsday_dispatch_html, send_resend_email, clean_val

def resolve_neon_ip(hostname):
    try:
        req = urllib.request.Request(
            f"https://1.1.1.1/dns-query?name={hostname}&type=A",
            headers={"accept": "application/dns-json"}
        )
        with urllib.request.urlopen(req, timeout=5) as resp:
            data = json.loads(resp.read().decode())
            for ans in data.get("Answer", []):
                if ans.get("type") == 1:
                    return ans["data"]
    except Exception as e:
        print(f"DoH resolution notice: {e}")
    return None

def connect_db():
    db_url = mm.DB_URL
    hostname = "ep-lively-boat-azbjpwf4-pooler.c-3.ap-southeast-1.aws.neon.tech"
    
    # 1. Try normal connection
    for attempt in range(2):
        try:
            return mm.get_db_connection()
        except Exception:
            time.sleep(1)
            
    # 2. Try with resolved hostaddr (bypasses Windows local DNS glitch)
    resolved_ip = resolve_neon_ip(hostname)
    if resolved_ip:
        print(f"Connecting via resolved IP ({resolved_ip})...")
        return psycopg2.connect(db_url, hostaddr=resolved_ip, connect_timeout=15)
        
    raise RuntimeError("Failed to connect to database after all retries")

def run_batch2_dispatch():
    print("=" * 65)
    print("   MOUNTAIN RUN — SPORTS DAY BATCH 2 MEDAL DISPATCH & DB UPDATE  ")
    print("=" * 65 + "\n")

    excel_path = "SPORTS_DAY_REMAINING_APPROVED_DISPATCH_LIST_UPDATED.xlsx"
    if not os.path.exists(excel_path):
        print(f"Error: {excel_path} not found!")
        return

    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Dispatch List"]

    runners = []
    for r in range(4, ws.max_row + 1):
        bib = clean_val(ws.cell(row=r, column=2).value)
        awb = clean_val(ws.cell(row=r, column=18).value)
        if bib and awb:
            name = clean_val(ws.cell(row=r, column=4).value) or clean_val(ws.cell(row=r, column=3).value)
            email = clean_val(ws.cell(row=r, column=6).value)
            dist = clean_val(ws.cell(row=r, column=7).value)
            courier = clean_val(ws.cell(row=r, column=19).value)
            link = clean_val(ws.cell(row=r, column=20).value)
            line1 = clean_val(ws.cell(row=r, column=10).value)
            line2 = clean_val(ws.cell(row=r, column=11).value)
            city = clean_val(ws.cell(row=r, column=12).value)
            state = clean_val(ws.cell(row=r, column=13).value)
            pincode = clean_val(ws.cell(row=r, column=14).value)

            runners.append({
                "row": r,
                "bib": bib,
                "name": name,
                "email": email,
                "distance": dist,
                "courier": courier or "Courier Partner",
                "awb": awb,
                "link": link or "https://mountainrun.in",
                "line1": line1,
                "line2": line2,
                "city": city,
                "state": state,
                "pincode": pincode
            })

    print(f"Found {len(runners)} athletes in Batch 2 with valid tracking details.\n")

    conn = connect_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)

    results = []
    sent_count = 0
    failed_count = 0

    for idx, r in enumerate(runners, start=1):
        payload = {
            "runnerName": r["name"],
            "bibNumber": r["bib"],
            "eventTitle": "Sports Day Celebration 2026",
            "distance": f"{r['distance']} Challenge" if r["distance"] else "Sports Day Finisher",
            "courier": r["courier"],
            "trackingNumber": r["awb"],
            "trackingUrl": r["link"],
            "shippingLine1": r["line1"],
            "shippingLine2": r["line2"],
            "shippingCity": r["city"],
            "shippingState": r["state"],
            "shippingPincode": r["pincode"],
        }

        subject = f"Your Finisher Medal is on the Way! 🏅 — {payload['eventTitle']}"
        html = build_sportsday_dispatch_html(payload)

        print(f"[{idx:2d}/{len(runners)}] Sending to {r['name']} <{r['email']}> | BIB: {r['bib']} | {r['courier']}: {r['awb']}...", end=" ", flush=True)

        try:
            # 1. Send Resend Email
            res = send_resend_email(r["email"], subject, html)
            res_id = res.get("id")
            sent_count += 1
            print(f"✓ Sent! (ID: {res_id})", end=" ", flush=True)

            # 2. Update Database
            cur.execute("""
                INSERT INTO "MedalDelivery" ("id", "registrationId", "status", "courier", "trackingNumber", "dispatchedAt")
                SELECT gen_random_uuid(), r.id, 'DISPATCHED', %s, %s, NOW()
                FROM "Registration" r
                WHERE r."bibNumber" = %s
                ON CONFLICT ("registrationId") DO UPDATE
                SET "status" = 'DISPATCHED', 
                    "courier" = EXCLUDED."courier", 
                    "trackingNumber" = EXCLUDED."trackingNumber", 
                    "dispatchedAt" = NOW()
            """, (r["courier"], r["awb"], r["bib"]))
            conn.commit()
            print("✓ DB Updated!", flush=True)

            # 3. Update Excel cell
            ws.cell(row=r["row"], column=17).value = "DISPATCHED"

            results.append({
                "bib": r["bib"],
                "name": r["name"],
                "email": r["email"],
                "awb": r["awb"],
                "courier": r["courier"],
                "resendId": res_id,
                "dbStatus": "DISPATCHED",
                "status": "SUCCESS"
            })

        except Exception as ex:
            print(f"❌ FAILED: {ex}", flush=True)
            failed_count += 1
            if conn:
                conn.rollback()
            results.append({
                "bib": r["bib"],
                "name": r["name"],
                "email": r["email"],
                "awb": r["awb"],
                "courier": r["courier"],
                "error": str(ex),
                "status": "FAILED"
            })

        time.sleep(0.8)

    # Save updated Excel
    try:
        wb.save(excel_path)
        print(f"\n✓ Updated Excel saved: {excel_path}")
    except Exception as e:
        print(f"Could not save Excel: {e}")

    # Query latest DB counts
    cur.execute("""
        SELECT 
            COUNT(*) as total_approved,
            COUNT(CASE WHEN m.status = 'DISPATCHED' THEN 1 END) as already_dispatched,
            COUNT(CASE WHEN m.status IS NULL OR m.status = 'PENDING' THEN 1 END) as pending_dispatch
        FROM "Registration" r
        JOIN "Event" e ON r."eventId" = e.id
        LEFT JOIN "MedalDelivery" m ON m."registrationId" = r.id
        WHERE e.slug = 'sports-day-celebration' AND r."proofStatus" = 'APPROVED'
    """)
    stats = cur.fetchone()
    conn.close()

    log_file = "sportsday_batch2_dispatch_results.json"
    with open(log_file, "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2, ensure_ascii=False)

    print("\n" + "=" * 65)
    print(f"   BATCH 2 COMPLETE: Sent: {sent_count} | Failed: {failed_count}")
    print(f"   DB STATS (SPORTS DAY APPROVED):")
    print(f"     Total Approved   : {stats['total_approved']}")
    print(f"     Dispatched       : {stats['already_dispatched']}")
    print(f"     Pending Dispatch : {stats['pending_dispatch']}")
    print("=" * 65)

if __name__ == "__main__":
    run_batch2_dispatch()
