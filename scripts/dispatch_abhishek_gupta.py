import os
import sys
import json
import time
from psycopg2.extras import RealDictCursor
import openpyxl

sys.path.insert(0, os.path.abspath("."))
import scripts.dispatch_sportsday_batch2 as db2
from scripts.dispatch_sportsday_tracking_emails import build_sportsday_dispatch_html, send_resend_email

def dispatch_abhishek():
    bib = "SDC-110225"
    awb = "100468115408"
    courier = "Trackon"
    tracking_url = "https://www.trackon.in/courier-tracking"

    print("=" * 60)
    print(f"   DISPATCHING FINISHER MEDAL FOR {bib} (Abhishek Gupta)")
    print("=" * 60 + "\n")

    conn = db2.connect_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)

    cur.execute("""
        SELECT 
            r.id as registration_id,
            r."bibNumber" as bib_number,
            u.name as user_name,
            u.email,
            COALESCE(r."shippingPhone", u.phone) as phone,
            r.distance,
            r."shippingName" as shipping_name,
            r."shippingLine1" as address_line1,
            COALESCE(r."shippingLine2", '') as address_line2,
            r."shippingCity" as city,
            r."shippingState" as state,
            r."shippingPincode" as pincode,
            r."proofStatus" as proof_status,
            e.title as event_title,
            m.status as medal_status
        FROM "Registration" r
        JOIN "User" u ON r."userId" = u.id
        JOIN "Event" e ON r."eventId" = e.id
        LEFT JOIN "MedalDelivery" m ON m."registrationId" = r.id
        WHERE r."bibNumber" = %s
    """, (bib,))

    runner = cur.fetchone()
    if not runner:
        print(f"❌ Athlete with BIB {bib} not found in database!")
        conn.close()
        return

    name = runner["shipping_name"] or runner["user_name"]
    email = runner["email"]

    print(f"Athlete Details:")
    print(f"  Name    : {name}")
    print(f"  Email   : {email}")
    print(f"  Phone   : {runner['phone']}")
    print(f"  Event   : {runner['event_title']}")
    print(f"  Distance: {runner['distance']}")
    print(f"  Address : {runner['address_line1']}, {runner['city']}, {runner['state']} - {runner['pincode']}")
    print(f"  Courier : {courier}")
    print(f"  AWB     : {awb}")
    print(f"  TrackUrl: {tracking_url}\n")

    payload = {
        "runnerName": name,
        "bibNumber": bib,
        "eventTitle": runner["event_title"] or "Sports Day Celebration 2026",
        "distance": f"{runner['distance']} Challenge" if runner["distance"] else "Sports Day Finisher",
        "courier": courier,
        "trackingNumber": awb,
        "trackingUrl": tracking_url,
        "shippingLine1": runner["address_line1"],
        "shippingLine2": runner["address_line2"],
        "shippingCity": runner["city"],
        "shippingState": runner["state"],
        "shippingPincode": runner["pincode"],
    }

    subject = f"Your Finisher Medal is on the Way! 🏅 — {payload['eventTitle']}"
    html = build_sportsday_dispatch_html(payload)

    print(f"Sending dispatch email to {name} <{email}>...", end=" ", flush=True)

    try:
        # 1. Send email via Resend
        res = send_resend_email(email, subject, html)
        res_id = res.get("id")
        print(f"✓ Sent successfully! (Resend ID: {res_id})", flush=True)

        # 2. Update Database
        cur.execute("""
            INSERT INTO "MedalDelivery" ("id", "registrationId", "status", "courier", "trackingNumber", "dispatchedAt")
            SELECT gen_random_uuid(), r.id, 'DISPATCHED', %s, %s, NOW()
            FROM "Registration" r
            WHERE r."bibNumber" = %s
            ON CONFLICT ("registrationId") DO UPDATE
            SET "status" = 'DISPATCHED', 
                "courier" = EXCLUDED."courier", 
                "trackingNumber" = EXCLUDED."trackingNumber", 
                "dispatchedAt" = NOW()
        """, (courier, awb, bib))
        conn.commit()
        print("✓ Database MedalDelivery updated to DISPATCHED!", flush=True)

    except Exception as ex:
        print(f"\n❌ Error during dispatch: {ex}")
        conn.rollback()
        conn.close()
        return

    conn.close()

    # 3. Update Excel sheets if present
    for excel_file in ["SPORTS_DAY_DISPATCH_LIST_v2.xlsx", "tracking id list.xlsx"]:
        if os.path.exists(excel_file):
            try:
                wb = openpyxl.load_workbook(excel_file)
                if "Dispatch List" in wb.sheetnames:
                    ws = wb["Dispatch List"]
                    for row in range(4, ws.max_row + 1):
                        if str(ws.cell(row=row, column=2).value).strip() == bib:
                            ws.cell(row=row, column=17).value = "DISPATCHED"
                            ws.cell(row=row, column=18).value = awb
                            ws.cell(row=row, column=19).value = courier
                            print(f"✓ Updated row {row} in {excel_file}")
                    wb.save(excel_file)
            except Exception as e:
                print(f"⚠️ Could not update {excel_file}: {e}")

    print("\n" + "=" * 60)
    print("   DISPATCH COMPLETE FOR ABHISHEK GUPTA!")
    print("=" * 60)

if __name__ == "__main__":
    dispatch_abhishek()
