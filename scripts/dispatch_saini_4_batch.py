import os
import sys
import json
import time
import datetime
import urllib.request
import urllib.error
import openpyxl
import psycopg2
from psycopg2.extras import RealDictCursor

sys.path.insert(0, os.path.abspath("."))
import scripts.medal_dispatch_manager as mm

ENV = mm.get_backend_env()
RESEND_API_KEY = ENV.get("RESEND_API_KEY")
RESEND_FROM = ENV.get("RESEND_FROM_EMAIL", "Mountain Run <onboarding@mountainrun.in>")
DB_URL = ENV.get("DATABASE_URL")

def get_db_connection(max_retries=5):
    for attempt in range(max_retries):
        try:
            return psycopg2.connect(DB_URL, connect_timeout=15)
        except Exception as e:
            if attempt < max_retries - 1:
                time.sleep(2)
            else:
                raise e

def get_courier_details(raw_courier, tracking_number):
    c_upper = str(raw_courier or "").upper().strip()
    t_str = str(tracking_number or "").strip()

    if any(k in c_upper for k in ["DELHIV", "DELHIVAR", "DELHIVERY", "DALEVER"]) or t_str.startswith("428"):
        return {
            "name": "Delhivery Express",
            "url": f"https://www.delhivery.com/track/package/{t_str}"
        }
    elif any(k in c_upper for k in ["BLUEDART", "BLUE DART"]):
        return {
            "name": "Blue Dart Express",
            "url": "https://www.bluedart.com/tracking"
        }
    else:
        return {
            "name": "DTDC Express",
            "url": "https://www.dtdc.com/track-your-shipment/"
        }

def send_resend_email(to_email, subject, html_content, max_retries=4):
    if not RESEND_API_KEY:
        raise ValueError("RESEND_API_KEY not found in backend/.env")
    url = "https://api.resend.com/emails"
    headers = {
        "Authorization": f"Bearer {RESEND_API_KEY}",
        "Content-Type": "application/json",
        "User-Agent": "MountainRunDispatch/1.0",
    }
    payload = {
        "from": RESEND_FROM,
        "to": [to_email],
        "subject": subject,
        "html": html_content,
        "reply_to": "mountainrunofficial@gmail.com",
    }
    for attempt in range(max_retries):
        try:
            req = urllib.request.Request(url, data=json.dumps(payload).encode("utf-8"), headers=headers, method="POST")
            with urllib.request.urlopen(req, timeout=20) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            if e.code == 429 and attempt < max_retries - 1:
                wait_s = (attempt + 1) * 3
                print(f"       ⏳ [429 Rate-Limit] Waiting {wait_s}s...")
                time.sleep(wait_s)
                continue
            err_msg = e.read().decode("utf-8") if e.fp else str(e)
            raise RuntimeError(f"HTTP {e.code}: {err_msg}")
        except (urllib.error.URLError, TimeoutError) as e:
            if attempt < max_retries - 1:
                wait_s = (attempt + 1) * 3
                print(f"       ⏳ [Network Glitch] Retrying in {wait_s}s ({e})...")
                time.sleep(wait_s)
                continue
            raise RuntimeError(f"Network error: {e}")

def process_saini_4(live=False):
    excel_path = os.path.join(os.getcwd(), "SAINI4.xlsx")
    if not os.path.exists(excel_path):
        print(f"Error: {excel_path} not found")
        return

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb["Courier Dispatch Ready"]

    conn = get_db_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute("""
        SELECT 
            r.id as registration_id,
            r."bibNumber" as bib,
            r.distance,
            r."shippingName" as shipping_name,
            r."shippingPhone" as shipping_phone,
            r."shippingLine1" as address_line1,
            COALESCE(r."shippingLine2", '') as address_line2,
            r."shippingCity" as city,
            r."shippingState" as state,
            r."shippingPincode" as pincode,
            u.name as user_name,
            u.email as user_email,
            e.title as event_title,
            m.id as medal_delivery_id,
            m.status as medal_status,
            m."trackingNumber" as db_tracking_number,
            m.courier as db_courier,
            m."dispatchedAt" as db_dispatched_at
        FROM "Registration" r
        JOIN "User" u ON r."userId" = u.id
        JOIN "Event" e ON r."eventId" = e.id
        LEFT JOIN "MedalDelivery" m ON m."registrationId" = r.id
        WHERE r."bibNumber" ILIKE 'RAR-%'
    """)
    db_rows = cur.fetchall()
    db_by_bib = {r['bib'].strip().upper(): r for r in db_rows if r['bib']}

    already_done = []
    to_update = []
    skipped_no_tracking = []

    for r in range(4, sheet.max_row + 1):
        sno = sheet.cell(r, 1).value
        bib = str(sheet.cell(r, 2).value or '').strip()
        name = str(sheet.cell(r, 3).value or '').strip()
        phone = str(sheet.cell(r, 4).value or '').strip()
        email = str(sheet.cell(r, 6).value or '').strip()
        # Col 16 is AVM (AWB Tracking ID)
        tracking = str(sheet.cell(r, 16).value or '').strip()
        courier_raw = "DTDC"

        if not bib and not name:
            continue

        if not tracking or tracking.upper() in ['NONE', 'NAN', 'NULL', '-', '', 'NO']:
            skipped_no_tracking.append({
                'sno': sno, 'row': r, 'bib': bib, 'name': name, 'raw_tracking': tracking
            })
            continue

        d = db_by_bib.get(bib.upper())
        if not d:
            print(f"⚠️ Warning: BIB {bib} not found in DB!")
            continue

        courier_info = get_courier_details(courier_raw, tracking)
        runner_name = d["shipping_name"] or d["user_name"] or name
        to_email = d["user_email"] or email
        dist_str = f"{d['distance']} Challenge" if d['distance'] else "Finisher"

        record_info = {
            'sno': sno,
            'row': r,
            'bib': bib,
            'runner_name': runner_name,
            'to_email': to_email,
            'tracking': tracking,
            'courier_info': courier_info,
            'event_title': d['event_title'] or "Rise and Run Virtual Challenge 2026",
            'distance': dist_str,
            'db_reg': d,
            'db_tracking': (d['db_tracking_number'] or '').strip(),
            'db_status': d['medal_status']
        }

        if record_info['db_tracking'] and record_info['db_status'] == 'DISPATCHED':
            already_done.append(record_info)
        else:
            to_update.append(record_info)

    print(f"\n==========================================")
    print(f" SAINI4.xlsx Dispatch Processing")
    print(f" Mode: {'LIVE (SENDING EMAILS & UPDATING DB)' if live else 'DRY-RUN (SIMULATION ONLY)'}")
    print(f"==========================================")
    print(f"Total Athletes Checked in Sheet : {len(already_done) + len(to_update) + len(skipped_no_tracking)}")
    print(f"Athletes to Update & Email      : {len(to_update)}")
    print(f"Athletes Without Tracking (Skip): {len(skipped_no_tracking)}")
    print(f"Already Dispatched in DB (Skip) : {len(already_done)}\n")

    if skipped_no_tracking:
        print("Skipped Athletes (No Tracking):")
        for s in skipped_no_tracking:
            print(f"  - S.No {s['sno']} | BIB: {s['bib']} | Name: {s['name']} (Tracking field: '{s['raw_tracking']}')")
        print()

    now_iso = datetime.datetime.now(datetime.timezone.utc)
    results = []

    for idx, item in enumerate(to_update, 1):
        bib = item["bib"]
        runner = item["runner_name"]
        email = item["to_email"]
        courier_name = item["courier_info"]["name"]
        courier_url = item["courier_info"]["url"]
        tracking = item["tracking"]
        db_reg = item["db_reg"]

        payload = {
            "to": email,
            "runnerName": runner,
            "bibNumber": bib,
            "eventTitle": item["event_title"],
            "distance": item["distance"],
            "courier": courier_name,
            "trackingNumber": tracking,
            "trackingUrl": courier_url,
            "shippingLine1": db_reg["address_line1"],
            "shippingLine2": db_reg["address_line2"],
            "shippingCity": db_reg["city"],
            "shippingState": db_reg["state"],
            "shippingPincode": db_reg["pincode"],
        }

        if not live:
            print(f"[{idx:02d}/{len(to_update)}] [DRY-RUN] Would send to: {runner} <{email}> | {courier_name} {tracking}")
            results.append({
                "bib": bib, "runner": runner, "email": email, "courier": courier_name,
                "tracking": tracking, "status": "DRY_RUN_OK"
            })
        else:
            subject = f"Your Finisher Medal is on the Way! 🏅 — {payload['eventTitle']}"
            html = mm.build_medal_dispatch_html(payload)
            print(f"[{idx:02d}/{len(to_update)}] Sending to {runner} <{email}> [{courier_name} {tracking}]...", end=" ", flush=True)
            try:
                res = send_resend_email(email, subject, html)
                res_id = res.get("id", "OK")

                # Ensure DB connection is alive
                try:
                    if conn.closed != 0:
                        conn = get_db_connection()
                        cur = conn.cursor(cursor_factory=RealDictCursor)
                except Exception:
                    conn = get_db_connection()
                    cur = conn.cursor(cursor_factory=RealDictCursor)

                # Update DB
                cur.execute("""
                    INSERT INTO "MedalDelivery" ("id", "registrationId", "status", "courier", "trackingNumber", "trackingUrl", "dispatchedAt")
                    VALUES (gen_random_uuid(), %s, 'DISPATCHED', %s, %s, %s, %s)
                    ON CONFLICT ("registrationId") DO UPDATE
                    SET "status" = 'DISPATCHED',
                        "courier" = EXCLUDED."courier",
                        "trackingNumber" = EXCLUDED."trackingNumber",
                        "trackingUrl" = EXCLUDED."trackingUrl",
                        "dispatchedAt" = EXCLUDED."dispatchedAt"
                """, (db_reg["registration_id"], courier_name, tracking, courier_url, now_iso))
                conn.commit()

                print(f"✅ SENT! (ID: {res_id})", flush=True)
                results.append({
                    "bib": bib, "runner": runner, "email": email, "courier": courier_name,
                    "tracking": tracking, "status": "SENT", "resend_id": res_id
                })
                time.sleep(0.5)
            except Exception as e:
                try:
                    conn.rollback()
                except Exception:
                    pass
                print(f"❌ FAILED: {e}", flush=True)
                results.append({
                    "bib": bib, "runner": runner, "email": email, "courier": courier_name,
                    "tracking": tracking, "status": "FAILED", "error": str(e)
                })

    try:
        conn.close()
    except Exception:
        pass

    # Save log
    out_file = os.path.join(os.getcwd(), f"saini4_dispatch_{'live' if live else 'dryrun'}_log.json")
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2)
    print(f"\nLogged execution details to {out_file}")

    return results

if __name__ == "__main__":
    is_live = "--live" in sys.argv
    process_saini_4(live=is_live)
