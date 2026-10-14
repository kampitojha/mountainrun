import os
import sys
import json
import openpyxl
import psycopg2
from psycopg2.extras import RealDictCursor
from tshirt_list_manager import get_backend_env

def analyze():
    excel_path = os.path.join(os.getcwd(), 'SAINI.xlsx')
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    sheet = wb['Courier Dispatch Ready']
    sheet_records = []
    for r in range(4, sheet.max_row + 1):
        sno = sheet.cell(r, 1).value
        bib = sheet.cell(r, 2).value
        name = sheet.cell(r, 3).value
        phone = sheet.cell(r, 4).value
        email = sheet.cell(r, 6).value
        tracking_val = sheet.cell(r, 16).value
        wait_val = sheet.cell(r, 17).value
        courier_val = sheet.cell(r, 18).value
        delivery_status_val = sheet.cell(r, 19).value
        amount_val = sheet.cell(r, 20).value
        
        if not bib and not name:
            continue
            
        tracking_clean = str(tracking_val).strip() if tracking_val is not None else ""
        if tracking_clean.lower() in ['none', 'null', 'nan', '-', '']:
            tracking_clean = ""
            
        sheet_records.append({
            'row_idx': r,
            'sno': sno,
            'bib': str(bib).strip() if bib else "",
            'name': str(name).strip() if name else "",
            'phone': str(phone).strip() if phone else "",
            'email': str(email).strip() if email else "",
            'tracking_id': tracking_clean,
            'courier': str(courier_val).strip() if courier_val is not None else "",
            'delivery_status': str(delivery_status_val).strip() if delivery_status_val is not None else ""
        })
        
    print(f"Total athletes in sheet: {len(sheet_records)}")
    with_tracking = [rec for rec in sheet_records if rec['tracking_id']]
    without_tracking = [rec for rec in sheet_records if not rec['tracking_id']]
    print(f"Rows WITH Tracking ID in SAINI.xlsx: {len(with_tracking)}")
    print(f"Rows WITHOUT Tracking ID in SAINI.xlsx: {len(without_tracking)}")
    
    # DB check
    env = get_backend_env()
    conn = None
    import time
    for attempt in range(5):
        try:
            conn = psycopg2.connect(env['DATABASE_URL'], connect_timeout=10)
            break
        except Exception as e:
            time.sleep(2)
    if not conn:
        raise RuntimeError("Could not connect to DB after 5 attempts")
    cur = conn.cursor(cursor_factory=RealDictCursor)
    
    cur.execute('''
        SELECT 
            r.id as registration_id,
            r."bibNumber" as bib,
            r."shippingName" as shipping_name,
            u.name as user_name,
            u.email as user_email,
            r."shippingPhone" as shipping_phone,
            u.phone as user_phone,
            r.status as reg_status,
            r."proofStatus" as proof_status,
            e.id as event_id,
            e.title as event_title,
            e.slug as event_slug,
            m.id as medal_delivery_id,
            m.status as medal_status,
            m."trackingNumber" as db_tracking_number,
            m.courier as db_courier
        FROM "Registration" r
        JOIN "Event" e ON r."eventId" = e.id
        JOIN "User" u ON r."userId" = u.id
        LEFT JOIN "MedalDelivery" m ON m."registrationId" = r.id
    ''')
    db_regs = cur.fetchall()
    print(f"Total registrations in DB: {len(db_regs)}")
    
    rar_db_regs = [r for r in db_regs if 'rise' in r['event_title'].lower() or 'rise' in r['event_slug'].lower() or 'rar' in str(r['bib']).lower()]
    print(f"Rise and Run registrations in DB: {len(rar_db_regs)}")
    already_dispatched_in_db = [r for r in rar_db_regs if r['medal_status'] == 'DISPATCHED']
    already_tracking_in_db = [r for r in rar_db_regs if r['db_tracking_number']]
    print(f"Already DISPATCHED in DB: {len(already_dispatched_in_db)}")
    print(f"Already has trackingNumber in DB: {len(already_tracking_in_db)}")
    
    db_by_bib = {r['bib'].strip().upper(): r for r in db_regs if r['bib']}
    db_rar_by_bib = {r['bib'].strip().upper(): r for r in rar_db_regs if r['bib']}
    
    # Check sheet records with tracking matching to DB
    matched_tracking = []
    unmatched_tracking = []
    for rec in with_tracking:
        b = rec['bib'].upper()
        if b in db_rar_by_bib:
            matched_tracking.append({'sheet': rec, 'db': db_rar_by_bib[b], 'is_rar': True})
        elif b in db_by_bib:
            matched_tracking.append({'sheet': rec, 'db': db_by_bib[b], 'is_rar': False})
        else:
            unmatched_tracking.append(rec)
            
    print(f"\n--- MATCH RESULTS FOR RECORDS WITH TRACKING ID ---")
    print(f"Total with tracking in sheet: {len(with_tracking)}")
    print(f"Matched in Rise and Run DB by BIB: {sum(1 for m in matched_tracking if m['is_rar'])}")
    print(f"Matched in other events in DB: {sum(1 for m in matched_tracking if not m['is_rar'])}")
    print(f"Unmatched in DB: {len(unmatched_tracking)}")
    
    if unmatched_tracking:
        print("Unmatched entries:")
        for u in unmatched_tracking:
            print(f"  BIB={u['bib']}, Name={u['name']}, Tracking={u['tracking_id']}")

    # Let us see how many of the 117 total match Rise and Run in DB
    all_matched = sum(1 for rec in sheet_records if rec['bib'].upper() in db_rar_by_bib)
    print(f"\nOut of all {len(sheet_records)} in SAINI.xlsx, matched in Rise and Run DB: {all_matched}")

    # Check tracking numbers already in DB vs what's in sheet
    already_in_db = []
    new_in_sheet = []
    diff_in_db = []
    for m in matched_tracking:
        db_track = (m['db']['db_tracking_number'] or '').strip()
        sheet_track = m['sheet']['tracking_id'].strip()
        if db_track:
            if db_track == sheet_track:
                already_in_db.append(m)
            else:
                diff_in_db.append({'bib': m['sheet']['bib'], 'name': m['sheet']['name'], 'sheet_track': sheet_track, 'db_track': db_track})
        else:
            new_in_sheet.append(m)
            
    print(f"\nDB Tracking Number status for matched records:")
    print(f"  - DB already has the exact same tracking ID: {len(already_in_db)}")
    print(f"  - DB has NO tracking number yet (new from sheet): {len(new_in_sheet)}")
    print(f"  - DB has a DIFFERENT tracking number: {len(diff_in_db)}")
    if diff_in_db:
        for d in diff_in_db:
            print(f"    Diff: BIB={d['bib']}, Name={d['name']}, Sheet={d['sheet_track']}, DB={d['db_track']}")

    # Let's inspect the ones without tracking ID in sheet
    print(f"\n--- LIST OF ATHLETES WITHOUT TRACKING ID IN SAINI.xlsx ({len(without_tracking)}) ---")
    for rec in without_tracking:
        db_info = db_rar_by_bib.get(rec['bib'].upper())
        db_track = (db_info['db_tracking_number'] or '') if db_info else 'N/A'
        print(f"  S.No {rec['sno']}: BIB={rec['bib']}, Name={rec['name']}, Phone={rec['phone']}, City={rec.get('city')}, DB Tracking={db_track}")

    # Also check if any duplicate tracking IDs exist in the sheet
    tracking_ids = [rec['tracking_id'] for rec in with_tracking]
    import collections
    counts = collections.Counter(tracking_ids)
    duplicates = {k: v for k, v in counts.items() if v > 1}
    if duplicates:
        print(f"\nDuplicate tracking IDs in sheet: {duplicates}")
    else:
        print(f"\nNo duplicate tracking IDs found in sheet (all unique).")

    conn.close()

if __name__ == '__main__':
    analyze()
