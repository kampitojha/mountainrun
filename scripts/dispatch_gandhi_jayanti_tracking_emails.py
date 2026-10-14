import os
import sys
import json
import time
import datetime
import urllib.request
import urllib.error
import openpyxl
import psycopg2
from psycopg2.extras import RealDictCursor

sys.path.insert(0, os.path.abspath("."))
import scripts.medal_dispatch_manager as mm

ENV = mm.get_backend_env()
RESEND_API_KEY = ENV.get("RESEND_API_KEY")
RESEND_FROM = ENV.get("RESEND_FROM_EMAIL", "Mountain Run <onboarding@mountainrun.in>")
DB_URL = ENV.get("DATABASE_URL")

def get_db_connection(max_retries=5):
    for attempt in range(max_retries):
        try:
            return psycopg2.connect(DB_URL, connect_timeout=15)
        except Exception as e:
            if attempt < max_retries - 1:
                time.sleep(2)
            else:
                raise e

def get_courier_details(raw_courier, tracking_number):
    c_upper = str(raw_courier or "").upper().strip()
    t_str = str(tracking_number or "").strip()

    if any(k in c_upper for k in ["DELHIV", "DELHIVAR", "DELHIVERY", "DALEVER"]) or t_str.startswith("428") or t_str.startswith("398"):
        return {
            "name": "Delhivery Express",
            "url": f"https://www.delhivery.com/track/package/{t_str}"
        }
    elif any(k in c_upper for k in ["SHADOWFAX"]) or t_str.startswith("SF"):
        return {
            "name": "Shadowfax",
            "url": "https://tracker.shadowfax.in/#/"
        }
    elif any(k in c_upper for k in ["XPRESSBEES"]):
        return {
            "name": "Xpressbees",
            "url": "https://www.xpressbees.com/shipment/tracking"
        }
    elif any(k in c_upper for k in ["BLUEDART", "BLUE DART"]):
        return {
            "name": "Blue Dart Express",
            "url": "https://www.bluedart.com/tracking"
        }
    else:
        return {
            "name": "DTDC Express",
            "url": "https://www.dtdc.com/track-your-shipment/"
        }

def send_resend_email(to_email, subject, html_content, max_retries=4):
    if not RESEND_API_KEY:
        raise ValueError("RESEND_API_KEY not found in backend/.env")
    url = "https://api.resend.com/emails"
    headers = {
        "Authorization": f"Bearer {RESEND_API_KEY}",
        "Content-Type": "application/json",
        "User-Agent": "MountainRunDispatch/1.0",
    }
    payload = {
        "from": RESEND_FROM,
        "to": [to_email],
        "subject": subject,
        "html": html_content,
        "reply_to": "mountainrunofficial@gmail.com",
    }
    for attempt in range(max_retries):
        try:
            req = urllib.request.Request(url, data=json.dumps(payload).encode("utf-8"), headers=headers, method="POST")
            with urllib.request.urlopen(req, timeout=20) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            if e.code == 429 and attempt < max_retries - 1:
                wait_s = (attempt + 1) * 3
                print(f"       ⏳ [429 Rate-Limit] Waiting {wait_s}s...")
                time.sleep(wait_s)
                continue
            err_msg = e.read().decode("utf-8") if e.fp else str(e)
            raise RuntimeError(f"HTTP {e.code}: {err_msg}")
        except (urllib.error.URLError, TimeoutError) as e:
            if attempt < max_retries - 1:
                wait_s = (attempt + 1) * 3
                print(f"       ⏳ [Network Glitch] Retrying in {wait_s}s ({e})...")
                time.sleep(wait_s)
                continue
            raise RuntimeError(f"Network error: {e}")

def process_gandhi_jayanti_dispatch(live=False):
    excel_candidates = [
        os.path.join(os.getcwd(), "Complete gandhi jayanti.xlsx"),
        os.path.join(os.getcwd(), "GANDHI_JAYANTI_MEDAL_DISPATCH_LIST (1).xlsx")
    ]
    excel_path = None
    for cand in excel_candidates:
        if os.path.exists(cand):
            excel_path = cand
            break

    if not excel_path:
        print("Error: Excel file not found!")
        return

    print(f"Using Excel file: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb["Gandhi Jayanti 2026"]

    conn = get_db_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute("""
        SELECT 
            r.id as registration_id,
            r."bibNumber" as bib,
            r.distance,
            r."shippingName" as shipping_name,
            r."shippingPhone" as shipping_phone,
            r."shippingLine1" as address_line1,
            COALESCE(r."shippingLine2", '') as address_line2,
            r."shippingCity" as city,
            r."shippingState" as state,
            r."shippingPincode" as pincode,
            u.name as user_name,
            u.email as user_email,
            e.title as event_title,
            m.id as medal_delivery_id,
            m.status as medal_status,
            m."trackingNumber" as db_tracking_number,
            m.courier as db_courier,
            m."dispatchedAt" as db_dispatched_at
        FROM "Registration" r
        JOIN "User" u ON r."userId" = u.id
        JOIN "Event" e ON r."eventId" = e.id
        LEFT JOIN "MedalDelivery" m ON m."registrationId" = r.id
        WHERE r."bibNumber" ILIKE 'GJVR-%'
    """)
    db_rows = cur.fetchall()
    db_by_bib = {r['bib'].strip().upper(): r for r in db_rows if r['bib']}

    already_done = []
    to_update = []
    skipped_no_tracking = []
    not_in_db = []

    for r in range(4, sheet.max_row + 1):
        sno = sheet.cell(r, 1).value
        bib = str(sheet.cell(r, 2).value or '').strip()
        name = str(sheet.cell(r, 3).value or '').strip()
        phone = str(sheet.cell(r, 4).value or '').strip()
        email = str(sheet.cell(r, 5).value or '').strip()
        address = str(sheet.cell(r, 6).value or '').strip()
        city = str(sheet.cell(r, 7).value or '').strip()
        state = str(sheet.cell(r, 8).value or '').strip()
        pincode = str(sheet.cell(r, 9).value or '').strip()
        event_name = str(sheet.cell(r, 10).value or '').strip()
        dist_cat = str(sheet.cell(r, 11).value or '').strip()
        package_item = str(sheet.cell(r, 12).value or '').strip()
        tracking = str(sheet.cell(r, 13).value or '').strip()
        courier_raw = str(sheet.cell(r, 14).value or '').strip()

        if not bib and not name:
            continue

        if not tracking or tracking.upper() in ['NONE', 'NAN', 'NULL', '-', '', 'NO']:
            skipped_no_tracking.append({
                'sno': sno, 'row': r, 'bib': bib, 'name': name, 'raw_tracking': tracking
            })
            continue

        d = db_by_bib.get(bib.upper())
        courier_info = get_courier_details(courier_raw, tracking)

        if not d:
            not_in_db.append({
                'sno': sno, 'row': r, 'bib': bib, 'name': name, 'email': email,
                'tracking': tracking, 'phone': phone, 'address': address, 'city': city,
                'state': state, 'pincode': pincode, 'distance': dist_cat, 'event_title': event_name
            })
            continue

        runner_name = d["shipping_name"] or d["user_name"] or name
        to_email = email or d["user_email"]
        dist_str = f"{d['distance']} Run" if d['distance'] else dist_cat or "Finisher"

        record_info = {
            'sno': sno,
            'row': r,
            'bib': bib,
            'runner_name': runner_name,
            'to_email': to_email,
            'tracking': tracking,
            'courier_info': courier_info,
            'event_title': d['event_title'] or event_name or "Gandhi Jayanti Victory Run 2026",
            'distance': dist_str,
            'address_line1': d["address_line1"] or address,
            'address_line2': d["address_line2"] or "",
            'city': d["city"] or city,
            'state': d["state"] or state,
            'pincode': d["pincode"] or pincode,
            'db_reg': d,
            'db_tracking': (d['db_tracking_number'] or '').strip(),
            'db_status': d['medal_status']
        }

        # If already marked as DISPATCHED with tracking in DB, skip!
        if record_info['db_tracking'] and record_info['db_status'] == 'DISPATCHED':
            already_done.append(record_info)
        else:
            to_update.append(record_info)

    print(f"\n==========================================")
    print(f" Gandhi Jayanti Medal Dispatch Processing")
    print(f" Mode: {'LIVE (SENDING EMAILS & UPDATING DB)' if live else 'DRY-RUN (SIMULATION ONLY)'}")
    print(f"==========================================")
    print(f"Total Athletes in Sheet         : {len(already_done) + len(to_update) + len(skipped_no_tracking) + len(not_in_db)}")
    print(f"Athletes to Update & Email      : {len(to_update)}")
    print(f"Athletes Already Dispatched (Skip): {len(already_done)}")
    print(f"Athletes Not in DB (Warning)    : {len(not_in_db)}")
    print(f"Athletes Without Tracking (Skip): {len(skipped_no_tracking)}\n")

    if not_in_db:
        print("⚠️ Athletes with tracking not found in DB:")
        for n in not_in_db:
            print(f"  - S.No {n['sno']} | BIB: {n['bib']} | Name: {n['name']} | Tracking: {n['tracking']}")
        print()

    if already_done:
        print(f"ℹ️ Already Dispatched ({len(already_done)} athletes skipped):")
        for a in already_done[:5]:
            print(f"  - S.No {a['sno']} | BIB: {a['bib']} | Name: {a['runner_name']} | Tracking: {a['tracking']}")
        if len(already_done) > 5:
            print(f"  ... and {len(already_done) - 5} more.")
        print()

    now_iso = datetime.datetime.now(datetime.timezone.utc)
    results = []

    for idx, item in enumerate(to_update, 1):
        bib = item["bib"]
        runner = item["runner_name"]
        email = item["to_email"]
        courier_name = item["courier_info"]["name"]
        courier_url = item["courier_info"]["url"]
        tracking = item["tracking"]
        db_reg = item["db_reg"]

        payload = {
            "to": email,
            "runnerName": runner,
            "bibNumber": bib,
            "eventTitle": item["event_title"],
            "distance": item["distance"],
            "courier": courier_name,
            "trackingNumber": tracking,
            "trackingUrl": courier_url,
            "shippingLine1": item["address_line1"],
            "shippingLine2": item["address_line2"],
            "shippingCity": item["city"],
            "shippingState": item["state"],
            "shippingPincode": item["pincode"],
        }

        if not live:
            print(f"[{idx:02d}/{len(to_update)}] [DRY-RUN] Would send to: {runner} <{email}> | {courier_name} {tracking} ({courier_url})")
            results.append({
                "bib": bib, "runner": runner, "email": email, "courier": courier_name,
                "tracking": tracking, "tracking_url": courier_url, "status": "DRY_RUN_OK"
            })
        else:
            subject = f"Your Finisher Medal is on the Way! 🏅 — {payload['eventTitle']}"
            html = mm.build_medal_dispatch_html(payload)
            print(f"[{idx:02d}/{len(to_update)}] Sending to {runner} <{email}> [{courier_name} {tracking}]...", end=" ", flush=True)
            try:
                res = send_resend_email(email, subject, html)
                res_id = res.get("id", "OK")

                # Ensure DB connection is alive
                try:
                    if conn.closed != 0:
                        conn = get_db_connection()
                        cur = conn.cursor(cursor_factory=RealDictCursor)
                except Exception:
                    conn = get_db_connection()
                    cur = conn.cursor(cursor_factory=RealDictCursor)

                # Update DB
                cur.execute("""
                    INSERT INTO "MedalDelivery" ("id", "registrationId", "status", "courier", "trackingNumber", "trackingUrl", "dispatchedAt")
                    VALUES (gen_random_uuid(), %s, 'DISPATCHED', %s, %s, %s, %s)
                    ON CONFLICT ("registrationId") DO UPDATE
                    SET "status" = 'DISPATCHED',
                        "courier" = EXCLUDED."courier",
                        "trackingNumber" = EXCLUDED."trackingNumber",
                        "trackingUrl" = EXCLUDED."trackingUrl",
                        "dispatchedAt" = EXCLUDED."dispatchedAt"
                """, (db_reg["registration_id"], courier_name, tracking, courier_url, now_iso))
                conn.commit()

                print(f"✅ SENT! (ID: {res_id})", flush=True)
                results.append({
                    "bib": bib, "runner": runner, "email": email, "courier": courier_name,
                    "tracking": tracking, "status": "SENT", "resend_id": res_id
                })
                time.sleep(0.6)
            except Exception as e:
                try:
                    conn.rollback()
                except Exception:
                    pass
                print(f"❌ FAILED: {e}", flush=True)
                results.append({
                    "bib": bib, "runner": runner, "email": email, "courier": courier_name,
                    "tracking": tracking, "status": "FAILED", "error": str(e)
                })

    try:
        conn.close()
    except Exception:
        pass

    # Save log
    out_file = os.path.join(os.getcwd(), f"gandhi_jayanti_dispatch_{'live' if live else 'dryrun'}_log.json")
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2)
    print(f"\nLogged execution details to {out_file}")

    return results

if __name__ == "__main__":
    is_live = "--live" in sys.argv
    process_gandhi_jayanti_dispatch(live=is_live)
