import os
import sys
import psycopg2
from psycopg2.extras import RealDictCursor
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.abspath('scripts'))
from tshirt_list_manager import get_backend_env

conn = psycopg2.connect(get_backend_env()['DATABASE_URL'])
cur = conn.cursor(cursor_factory=RealDictCursor)

# 1. Fetch Gandhi Jayanti runners (141)
cur.execute('''
    SELECT 
        e.title as event_title,
        r."bibNumber" as bib_number,
        COALESCE(r."shippingName", u.name) as recipient_name,
        COALESCE(r."shippingPhone", u.phone) as phone,
        u.email,
        r."shippingLine1" as address_line1,
        COALESCE(r."shippingLine2", '') as address_line2,
        r."shippingCity" as city,
        r."shippingState" as state,
        r."shippingPincode" as pincode,
        r.distance as distance_category,
        r."finishTimeSeconds" as finish_time_seconds
    FROM "Registration" r
    JOIN "Event" e ON r."eventId" = e.id
    JOIN "User" u ON r."userId" = u.id
    WHERE e.slug = 'gandhi-jayanti-victory-run-2026'
      AND r."proofStatus" = 'APPROVED'
      AND (r."adminNote" IS NULL OR r."adminNote" NOT ILIKE '%%DUMMY%%')
      AND r."bibNumber" != 'GJVR-109452'
    ORDER BY r."registeredAt" ASC
''')
gandhi_runners = cur.fetchall()

# 2. Fetch other pending approved runners (38)
cur.execute('''
    SELECT 
        e.title as event_title,
        r."bibNumber" as bib_number,
        COALESCE(r."shippingName", u.name) as recipient_name,
        COALESCE(r."shippingPhone", u.phone) as phone,
        u.email,
        r."shippingLine1" as address_line1,
        COALESCE(r."shippingLine2", '') as address_line2,
        r."shippingCity" as city,
        r."shippingState" as state,
        r."shippingPincode" as pincode,
        r.distance as distance_category,
        r."finishTimeSeconds" as finish_time_seconds
    FROM "Registration" r
    JOIN "Event" e ON r."eventId" = e.id
    JOIN "User" u ON r."userId" = u.id
    LEFT JOIN "MedalDelivery" m ON r.id = m."registrationId"
    WHERE e.slug != 'gandhi-jayanti-victory-run-2026'
      AND r."proofStatus" = 'APPROVED'
      AND (m.status = 'PENDING' OR m.id IS NULL)
      AND (r."adminNote" IS NULL OR r."adminNote" NOT ILIKE '%%DUMMY%%')
    ORDER BY e.title, r."registeredAt" ASC
''')
other_runners = cur.fetchall()

conn.close()

print(f"Loaded Gandhi Jayanti: {len(gandhi_runners)} runners")
print(f"Loaded Other Events: {len(other_runners)} runners")
total_runners = len(gandhi_runners) + len(other_runners)
print(f"Total Combined: {total_runners} runners")

def clean_phone(p):
    if not p:
        return ""
    p_str = str(p).strip().replace(" ", "").replace("-", "")
    if p_str.startswith("+91"):
        p_str = p_str[3:]
    elif p_str.startswith("91") and len(p_str) > 10:
        p_str = p_str[2:]
    return p_str

def get_medal_desc(event_title, dist):
    dist_clean = dist.upper().replace(" ", "")
    if "Gandhi" in event_title:
        return f"Gandhi Jayanti Finisher Medal ({dist})"
    elif "Sports" in event_title:
        return f"Sports Day Finisher Medal ({dist})"
    elif "Rise" in event_title:
        return f"Rise & Run Finisher Medal ({dist})"
    elif "Independence" in event_title:
        return f"Independence Day Finisher Medal ({dist})"
    return f"Finisher Medal ({dist})"

def build_full_address(line1, line2, city, state, pin):
    parts = []
    if line1 and line1.strip():
        parts.append(line1.strip().rstrip(','))
    if line2 and line2.strip():
        parts.append(line2.strip().rstrip(','))
    if city and city.strip():
        parts.append(city.strip())
    if state and state.strip():
        parts.append(state.strip())
    addr = ", ".join(parts)
    if pin and str(pin).strip():
        addr += f" - {str(pin).strip()}"
    return addr

# Styles
font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
font_subtitle = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_data = Font(name="Calibri", size=10)
font_bold = Font(name="Calibri", size=10, bold=True)

fill_dark_green = PatternFill(start_color="0F5132", end_color="0F5132", fill_type="solid")
fill_header = PatternFill(start_color="198754", end_color="198754", fill_type="solid")
fill_alt = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")

thin_gray = Side(border_style="thin", color="D3D3D3")
border_cell = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")

headers = [
    "S.No",
    "BIB Number",
    "Recipient Name",
    "Primary Mobile (10-Digit)",
    "Email Address",
    "Full Shipping Address",
    "Address Line 1",
    "Address Line 2 / Landmark",
    "City",
    "State",
    "Pincode",
    "Event Name",
    "Distance Category",
    "Package Contents (Medal Item)"
]

wb = openpyxl.Workbook()
wb.remove(wb.active)  # remove default sheet

def populate_sheet(ws, sheet_title, runner_list, banner_title):
    ws.views.sheetView[0].showGridLines = True
    
    # Title Banner Row 1 & 2
    ws.merge_cells("A1:N1")
    title_cell = ws["A1"]
    title_cell.value = f"MOUNTAIN RUN  |  {banner_title.upper()}"
    title_cell.font = font_title
    title_cell.fill = fill_dark_green
    title_cell.alignment = align_center
    ws.row_dimensions[1].height = 28
    
    ws.merge_cells("A2:N2")
    sub_cell = ws["A2"]
    sub_cell.value = f"OFFICIAL APPROVED PROOFS — COURIER MEDAL DISPATCH LIST (TOTAL: {len(runner_list)} ATHLETES)"
    sub_cell.font = font_subtitle
    sub_cell.fill = fill_dark_green
    sub_cell.alignment = align_center
    ws.row_dimensions[2].height = 20

    # Headers Row 3
    ws.row_dimensions[3].height = 26
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_cell

    # Data Rows
    for row_idx, r in enumerate(runner_list, 1):
        actual_row = row_idx + 3
        ws.row_dimensions[actual_row].height = 22
        
        phone_clean = clean_phone(r['phone'])
        full_addr = build_full_address(r['address_line1'], r['address_line2'], r['city'], r['state'], r['pincode'])
        medal_item = get_medal_desc(r['event_title'], r['distance_category'])
        
        row_values = [
            row_idx,
            r['bib_number'],
            r['recipient_name'],
            phone_clean,
            r['email'],
            full_addr,
            r['address_line1'],
            r['address_line2'],
            r['city'],
            r['state'],
            str(r['pincode']).strip() if r['pincode'] else '',
            r['event_title'],
            r['distance_category'],
            medal_item
        ]
        
        is_even = (row_idx % 2 == 0)
        row_fill = fill_alt if is_even else None
        
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=actual_row, column=col_idx, value=val)
            cell.font = font_data
            cell.border = border_cell
            if row_fill:
                cell.fill = row_fill
            
            # Alignments
            if col_idx in (1, 2, 4, 11, 13):
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col[2:]: # ignore merged title rows
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
    
    # Specific column width overrides
    ws.column_dimensions['A'].width = 8   # S.No
    ws.column_dimensions['B'].width = 16  # Bib
    ws.column_dimensions['C'].width = 24  # Name
    ws.column_dimensions['D'].width = 18  # Phone
    ws.column_dimensions['E'].width = 28  # Email
    ws.column_dimensions['F'].width = 45  # Full Address
    ws.column_dimensions['G'].width = 30  # Addr 1
    ws.column_dimensions['H'].width = 20  # Addr 2
    ws.column_dimensions['I'].width = 18  # City
    ws.column_dimensions['J'].width = 18  # State
    ws.column_dimensions['K'].width = 12  # Pincode
    ws.column_dimensions['L'].width = 32  # Event
    ws.column_dimensions['M'].width = 16  # Distance
    ws.column_dimensions['N'].width = 35  # Package item

# Tab 1: All Combined Master (179)
all_combined = list(gandhi_runners) + list(other_runners)
ws_master = wb.create_sheet(title="All Events Master Dispatch")
populate_sheet(ws_master, "All Events Master Dispatch", all_combined, "ALL EVENTS MASTER MEDAL DISPATCH LIST")

# Tab 2: Gandhi Jayanti 2026 (141)
ws_gandhi = wb.create_sheet(title="Gandhi Jayanti 2026")
populate_sheet(ws_gandhi, "Gandhi Jayanti 2026", gandhi_runners, "GANDHI JAYANTI VICTORY RUN 2026")

# Tab 3: Sports Day & Backlog (38)
ws_other = wb.create_sheet(title="Other Events Backlog")
populate_sheet(ws_other, "Other Events Backlog", other_runners, "OTHER EVENTS PENDING MEDALS (SPORTS DAY, RAR, IDVR)")

# Tab 4: Summary & Distance Breakdown
ws_summary = wb.create_sheet(title="Summary & Counts")
ws_summary.views.sheetView[0].showGridLines = True

ws_summary["A1"] = "MOUNTAIN RUN MEDAL DISPATCH SUMMARY & INVENTORY BREAKDOWN"
ws_summary["A1"].font = font_title
ws_summary["A1"].fill = fill_dark_green
ws_summary.merge_cells("A1:E1")
ws_summary.row_dimensions[1].height = 28

sum_headers = ["Event Name", "Category / Distance", "Medal Count", "Sheet Tab"]
ws_summary.row_dimensions[3].height = 24
for c_idx, h in enumerate(sum_headers, 1):
    c = ws_summary.cell(row=3, column=c_idx, value=h)
    c.font = font_header
    c.fill = fill_header
    c.alignment = align_center

summary_rows = []

# Gandhi breakdown
g_dist = {}
for r in gandhi_runners:
    g_dist[r['distance_category']] = g_dist.get(r['distance_category'], 0) + 1
for d, cnt in sorted(g_dist.items()):
    summary_rows.append(["Gandhi Jayanti Victory Run 2026", d, cnt, "Gandhi Jayanti 2026"])

# Other breakdown
o_dist = {}
for r in other_runners:
    key = f"{r['event_title']} ({r['distance_category']})"
    o_dist[key] = o_dist.get(key, 0) + 1
for k, cnt in sorted(o_dist.items()):
    summary_rows.append([k.split(" (")[0], k.split(" (")[1].rstrip(")"), cnt, "Other Events Backlog"])

for r_idx, srow in enumerate(summary_rows, 4):
    for c_idx, val in enumerate(srow, 1):
        c = ws_summary.cell(row=r_idx, column=c_idx, value=val)
        c.font = font_data
        c.border = border_cell
        if c_idx in (2, 3):
            c.alignment = align_center
        else:
            c.alignment = align_left

tot_row = len(summary_rows) + 4
ws_summary.cell(row=tot_row, column=1, value="TOTAL MEDALS TO DISPATCH").font = font_bold
ws_summary.cell(row=tot_row, column=3, value=total_runners).font = font_bold
ws_summary.cell(row=tot_row, column=3).alignment = align_center
for c in range(1, 5):
    ws_summary.cell(row=tot_row, column=c).fill = fill_alt
    ws_summary.cell(row=tot_row, column=c).border = border_cell

ws_summary.column_dimensions['A'].width = 38
ws_summary.column_dimensions['B'].width = 22
ws_summary.column_dimensions['C'].width = 16
ws_summary.column_dimensions['D'].width = 28

output_path = os.path.join(os.getcwd(), "MOUNTAIN_RUN_MEDAL_DISPATCH_MASTER.xlsx")
wb.save(output_path)
print(f"\nSuccessfully generated Master Excel sheet at: {output_path}")

# Also save standalone Gandhi Jayanti sheet for single event use
gandhi_wb = openpyxl.Workbook()
gandhi_wb.remove(gandhi_wb.active)
ws_single = gandhi_wb.create_sheet(title="Gandhi Jayanti 2026")
populate_sheet(ws_single, "Gandhi Jayanti 2026", gandhi_runners, "GANDHI JAYANTI VICTORY RUN 2026")
gandhi_path = os.path.join(os.getcwd(), "GANDHI_JAYANTI_MEDAL_DISPATCH_LIST.xlsx")
gandhi_wb.save(gandhi_path)
print(f"Successfully generated standalone Gandhi Jayanti Excel sheet at: {gandhi_path}")
