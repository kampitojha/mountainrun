import sys
import os
sys.path.insert(0, os.path.abspath("."))
import scripts.medal_dispatch_manager as mm
from psycopg2.extras import RealDictCursor
import openpyxl

def sync_tshirt_sizes():
    excel_path = "SPORTS_DAY_LEADERBOARD_TSHIRT_WINNERS.xlsx"
    if not os.path.exists(excel_path):
        print(f"Error: {excel_path} not found.")
        return

    print("Connecting to database to fetch submitted T-Shirt sizes...")
    conn = mm.get_db_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)

    cur.execute("""
        SELECT r."bibNumber", r."tshirtSize", r."tshirtSubmittedAt", r."shippingName", u.name as user_name
        FROM "Registration" r
        JOIN "User" u ON r."userId" = u.id
        JOIN "Event" e ON r."eventId" = e.id
        WHERE e.slug = 'sports-day-celebration'
          AND r."tshirtSize" IS NOT NULL
    """)
    db_sizes = {r["bibNumber"].strip().upper(): r for r in cur.fetchall()}
    conn.close()

    print(f"Found {len(db_sizes)} athletes with submitted T-Shirt sizes in DB.")

    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Top 3 T-Shirt Winners"]

    updated_count = 0
    for row_idx in range(5, ws.max_row + 1):
        bib_cell = ws.cell(row=row_idx, column=4)
        if not bib_cell.value:
            continue
        bib = str(bib_cell.value).strip().upper()
        if bib in db_sizes:
            data = db_sizes[bib]
            size = data["tshirtSize"]
            # Column 13 is T-Shirt Size
            ws.cell(row=row_idx, column=13, value=size)
            # Column 14 is Dispatch Status
            ws.cell(row=row_idx, column=14, value=f"Size Confirmed ({size})")
            updated_count += 1
            print(f"✓ [{bib}] {data['shippingName'] or data['user_name']} -> Size: {size}")

    wb.save(excel_path)
    print(f"\n[DONE] Successfully synced {updated_count} sizes into {excel_path}!")

if __name__ == "__main__":
    sync_tshirt_sizes()
