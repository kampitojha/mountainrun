#!/usr/bin/env python3
"""
Sync Rise and Run T-Shirt Sizes from PostgreSQL to RISE_AND_RUN_LEADERBOARD_TSHIRT_WINNERS.xlsx
"""

import os
import sys
import openpyxl
from collections import Counter
import psycopg2
from psycopg2.extras import RealDictCursor

sys.path.insert(0, os.path.abspath("scripts"))
from tshirt_list_manager import get_backend_env

excel_path = "RISE_AND_RUN_LEADERBOARD_TSHIRT_WINNERS.xlsx"
if not os.path.exists(excel_path):
    print(f"Error: {excel_path} not found! Run scripts/create_rise_and_run_top3_excel.py first.")
    sys.exit(1)

wb = openpyxl.load_workbook(excel_path)
ws = wb["Top 3 T-Shirt Winners"]

excel_rows = {}
for r in range(5, ws.max_row + 1):
    bib = ws.cell(row=r, column=4).value
    if bib:
        bib_clean = str(bib).strip().upper()
        excel_rows[bib_clean] = {
            "row": r,
            "sno": ws.cell(row=r, column=1).value,
            "category": ws.cell(row=r, column=2).value,
            "rank": ws.cell(row=r, column=3).value,
            "name": ws.cell(row=r, column=5).value,
            "size": ws.cell(row=r, column=13).value,
            "status": ws.cell(row=r, column=14).value
        }

total_winners = len(excel_rows)
env = get_backend_env()
conn = psycopg2.connect(env.get("DATABASE_URL"), connect_timeout=10)
cur = conn.cursor(cursor_factory=RealDictCursor)

cur.execute("""
    SELECT 
        r."bibNumber", 
        r."tshirtSize", 
        r."tshirtSubmittedAt", 
        r."shippingName",
        r."shippingPhone",
        r."shippingLine1",
        r."shippingCity",
        r."shippingState",
        r."shippingPincode"
    FROM "Registration" r
    JOIN "Event" e ON r."eventId" = e.id
    WHERE e.slug = 'rise-and-run'
      AND r."bibNumber" = ANY(%s)
""", (list(excel_rows.keys()),))

db_rows = cur.fetchall()
conn.close()

db_data = {r["bibNumber"].strip().upper(): r for r in db_rows}

new_updates = []
all_current_sizes = {}

for bib, info in excel_rows.items():
    db_rec = db_data.get(bib)
    db_size = db_rec["tshirtSize"].strip().upper() if db_rec and db_rec["tshirtSize"] else None
    current_size = str(info["size"]).strip().upper() if info["size"] else None

    if db_size and db_size != current_size:
        new_updates.append({
            "bib": bib,
            "name": info["name"],
            "old_size": current_size,
            "new_size": db_size,
        })
        ws.cell(row=info["row"], column=13, value=db_size)
        ws.cell(row=info["row"], column=14, value=f"Confirmed ({db_size})")
        all_current_sizes[bib] = db_size
    elif current_size:
        all_current_sizes[bib] = current_size
    else:
        all_current_sizes[bib] = None

if new_updates:
    wb.save(excel_path)
    print(f"\n[UPDATED] Successfully synced {len(new_updates)} sizes into {excel_path}!")
else:
    print(f"\n[NO CHANGES] No new sizes were submitted in DB since last sync.")

confirmed_count = sum(1 for s in all_current_sizes.values() if s)
pending_count = total_winners - confirmed_count

print(f"\n================ SUMMARY ================")
print(f"Total Winners: {total_winners}")
print(f"Sizes Confirmed: {confirmed_count} / {total_winners}")
print(f"Sizes Pending: {pending_count} / {total_winners}")

print("\n--- ALL 15 ATHLETES STATUS ---")
for bib, info in excel_rows.items():
    s = all_current_sizes[bib]
    status_str = f"Confirmed ({s})" if s else "⏳ PENDING"
    print(f"BIB: {bib:12s} | {info['category']:7s} | {info['rank']:24s} | {info['name']:24s} | {status_str}")
