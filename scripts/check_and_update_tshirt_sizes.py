import sys
import os
sys.path.insert(0, os.path.abspath("."))
import time
import openpyxl
from collections import Counter
from psycopg2.extras import RealDictCursor
import scripts.medal_dispatch_manager as mm

excel_path = "SPORTS_DAY_LEADERBOARD_TSHIRT_WINNERS.xlsx"
if not os.path.exists(excel_path):
    print(f"Error: {excel_path} not found!")
    sys.exit(1)

wb = openpyxl.load_workbook(excel_path)
ws = wb["Top 3 T-Shirt Winners"]

excel_rows = {}
for r in range(5, ws.max_row + 1):
    bib = ws.cell(row=r, column=4).value
    if bib:
        bib_clean = str(bib).strip().upper()
        excel_rows[bib_clean] = {
            "row": r,
            "sno": ws.cell(row=r, column=1).value,
            "category": ws.cell(row=r, column=2).value,
            "rank": ws.cell(row=r, column=3).value,
            "name": ws.cell(row=r, column=5).value,
            "size": ws.cell(row=r, column=13).value,
            "status": ws.cell(row=r, column=14).value
        }

total_winners = len(excel_rows)
current_sizes_in_excel = {k: v for k, v in excel_rows.items() if v["size"]}
print(f"Total winners in Excel: {total_winners}")
print(f"Current sizes present in Excel: {len(current_sizes_in_excel)} / {total_winners}")

# Connect to DB
conn = None
for i in range(5):
    try:
        conn = mm.get_db_connection()
        break
    except Exception as e:
        print(f"DB retry {i+1}/5: {e}")
        time.sleep(2)

if not conn:
    raise RuntimeError("Failed to connect to database after 5 retries")

cur = conn.cursor(cursor_factory=RealDictCursor)

cur.execute("""
    SELECT 
        r."bibNumber", 
        r."tshirtSize", 
        r."tshirtSubmittedAt", 
        r."shippingName", 
        u.name as user_name,
        r."adminNote",
        r."shippingLine2",
        pu."reviewerNote"
    FROM "Registration" r
    JOIN "User" u ON r."userId" = u.id
    JOIN "Event" e ON r."eventId" = e.id
    LEFT JOIN "ProofUpload" pu ON pu."registrationId" = r.id
    WHERE e.slug = 'sports-day-celebration'
      AND r."bibNumber" = ANY(%s)
""", (list(excel_rows.keys()),))

db_rows = cur.fetchall()
conn.close()

db_data = {r["bibNumber"].strip().upper(): r for r in db_rows}

new_updates = []
all_current_sizes = {}

for bib, info in excel_rows.items():
    db_rec = db_data.get(bib)
    db_size = db_rec["tshirtSize"].strip().upper() if db_rec and db_rec["tshirtSize"] else None
    current_size = str(info["size"]).strip().upper() if info["size"] else None

    # Check if a new size was submitted
    if db_size and not current_size:
        new_updates.append({
            "bib": bib,
            "name": info["name"],
            "category": info["category"],
            "rank": info["rank"],
            "old_size": None,
            "new_size": db_size,
            "submitted_at": db_rec.get("tshirtSubmittedAt")
        })
        ws.cell(row=info["row"], column=13, value=db_size)
        ws.cell(row=info["row"], column=14, value=f"Size Confirmed ({db_size})")
        all_current_sizes[bib] = db_size
    elif db_size and current_size and db_size != current_size:
        new_updates.append({
            "bib": bib,
            "name": info["name"],
            "category": info["category"],
            "rank": info["rank"],
            "old_size": current_size,
            "new_size": db_size,
            "submitted_at": db_rec.get("tshirtSubmittedAt")
        })
        ws.cell(row=info["row"], column=13, value=db_size)
        ws.cell(row=info["row"], column=14, value=f"Size Confirmed ({db_size})")
        all_current_sizes[bib] = db_size
    elif current_size:
        all_current_sizes[bib] = current_size
    else:
        all_current_sizes[bib] = None

# Save if updated
if new_updates:
    wb.save(excel_path)
    print(f"\n[UPDATED] Successfully updated {len(new_updates)} new/changed sizes into {excel_path}!")
else:
    print(f"\n[NO CHANGES] No new sizes were submitted in DB since last sync.")

confirmed_count = sum(1 for s in all_current_sizes.values() if s)
pending_count = total_winners - confirmed_count

print(f"\n================ SUMMARY ================")
print(f"Total Winners: {total_winners}")
print(f"Sizes Confirmed: {confirmed_count} / {total_winners}")
print(f"Sizes Still Pending: {pending_count} / {total_winners}")

size_counts = Counter([s for s in all_current_sizes.values() if s])
print("\nSize breakdown:", dict(size_counts))

print("\n--- ALL 21 ATHLETES STATUS ---")
for bib, info in excel_rows.items():
    s = all_current_sizes[bib]
    status_str = f"Confirmed ({s})" if s else "⏳ PENDING"
    print(f"BIB: {bib:10s} | {info['category']:7s} | {info['rank']:10s} | {info['name']:24s} | {status_str}")
