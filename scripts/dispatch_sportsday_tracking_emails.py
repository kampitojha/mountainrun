import os
import sys
import json
import time
import datetime
import urllib.request
import urllib.error
import openpyxl
from psycopg2.extras import RealDictCursor

sys.path.insert(0, os.path.abspath("."))
import scripts.medal_dispatch_manager as mm

ENV = mm.get_backend_env()
RESEND_API_KEY = ENV.get("RESEND_API_KEY")
RESEND_FROM = ENV.get("RESEND_FROM_EMAIL", "Mountain Run <onboarding@mountainrun.in>")

def build_sportsday_dispatch_html(payload):
    DARK_GREEN = "#1a3a2e"
    GOLD = "#c9a227"
    CREAM = "#fcfaf5"
    MUTED = "#8a7a5a"
    LINE = "#e8dfc8"

    tracking_num = str(payload.get("trackingNumber", "")).strip()
    courier_name = payload.get("courier", "Delhivery Express")
    tracking_url = payload.get("trackingUrl") or "https://mountainrun.in"

    address_parts = [
        payload.get("shippingLine1"),
        payload.get("shippingLine2"),
        payload.get("shippingCity"),
        f"{payload.get('shippingState', '')} - {payload.get('shippingPincode', '')}".strip(" -")
    ]
    address_display = ", ".join([p for p in address_parts if p and p != "—"])

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8"/>
  <meta name="viewport" content="width=device-width,initial-scale=1.0"/>
  <title>Your Finisher Medal is on the Way! — Mountain Run</title>
</head>
<body style="margin:0;padding:0;background-color:#f0ede5;font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,Helvetica,Arial,sans-serif;-webkit-font-smoothing:antialiased;">
  <div style="display:none;font-size:1px;color:#f0ede5;line-height:1px;max-height:0px;max-width:0px;opacity:0;overflow:hidden;">
    Great news {payload['runnerName']}! Your official finisher medal for {payload['eventTitle']} has been dispatched via {courier_name}. Tracking: {tracking_num} 🏅
  </div>

  <table width="100%" cellpadding="0" cellspacing="0" border="0" style="background-color:#f0ede5;padding:32px 12px;">
    <tr>
      <td align="center">
        <table width="600" cellpadding="0" cellspacing="0" border="0" style="max-width:600px;width:100%;background-color:{CREAM};border-radius:16px;overflow:hidden;box-shadow:0 10px 30px rgba(26,58,46,0.12);border:1px solid {LINE};">

          <!-- Header -->
          <tr>
            <td style="background-color:{DARK_GREEN};padding:0;">
              <!-- Tricolor ribbon -->
              <table width="100%" cellpadding="0" cellspacing="0" border="0">
                <tr>
                  <td width="33.3%" height="4" style="background-color:#FF9933;"></td>
                  <td width="33.4%" height="4" style="background-color:#ffffff;"></td>
                  <td width="33.3%" height="4" style="background-color:#138808;"></td>
                </tr>
              </table>
              <table width="100%" cellpadding="0" cellspacing="0" border="0">
                <tr>
                  <td style="padding:26px 32px 20px;text-align:center;">
                    <div style="display:inline-block;background:rgba(255,255,255,0.08);border:1px solid rgba(201,162,39,0.4);border-radius:50px;padding:6px 20px;margin-bottom:8px;">
                      <span style="font-size:12px;font-weight:800;letter-spacing:0.22em;text-transform:uppercase;color:{GOLD};">⛰️ MOUNTAIN RUN</span>
                    </div>
                    <p style="margin:0;font-size:10px;letter-spacing:0.25em;text-transform:uppercase;color:rgba(255,255,255,0.7);">Official Event Logistics · Dispatch Update</p>
                  </td>
                </tr>
              </table>
            </td>
          </tr>

          <!-- Hero Greeting -->
          <tr>
            <td style="padding:36px 36px 20px;text-align:center;">
              <div style="display:inline-block;width:64px;height:64px;line-height:64px;background-color:#fbf6e8;border:2px solid {GOLD};border-radius:50%;font-size:32px;margin-bottom:14px;">
                🏅
              </div>
              <h1 style="margin:0 0 6px;font-size:24px;font-weight:800;color:{DARK_GREEN};font-family:Georgia,serif;letter-spacing:-0.01em;">
                Your Finisher Medal is on the Way!
              </h1>
              <p style="margin:0 0 20px;font-size:12px;font-weight:700;letter-spacing:0.15em;text-transform:uppercase;color:{GOLD};">
                ── {payload['eventTitle']} ──
              </p>
              <p style="margin:0;font-size:15px;color:#444444;line-height:1.7;text-align:left;">
                Hi <strong style="color:{DARK_GREEN};">{payload['runnerName']}</strong>,<br/><br/>
                Congratulations once again on your incredible spirit and achievement! We are delighted to inform you that your official physical <strong>Finisher Medal</strong> has been carefully packed and handed over to our courier partner.
              </p>
            </td>
          </tr>

          <!-- Tracking Card -->
          <tr>
            <td style="padding:0 36px 24px;">
              <table width="100%" cellpadding="0" cellspacing="0" border="0" style="background-color:#ffffff;border:1.5px solid {LINE};border-radius:12px;overflow:hidden;box-shadow:0 4px 12px rgba(0,0,0,0.03);">
                <tr>
                  <td colspan="2" style="background-color:#f6f1e5;padding:12px 18px;border-bottom:1px solid {LINE};">
                    <table width="100%" cellpadding="0" cellspacing="0" border="0">
                      <tr>
                        <td align="left">
                          <span style="font-size:11px;font-weight:800;letter-spacing:0.12em;text-transform:uppercase;color:{DARK_GREEN};">📦 Consignment Details</span>
                        </td>
                        <td align="right">
                          <span style="display:inline-block;background-color:#e6f4ea;color:#137333;font-size:10px;font-weight:800;letter-spacing:0.08em;text-transform:uppercase;padding:3px 10px;border-radius:20px;border:1px solid #b7e1cd;">
                            ● Dispatched
                          </span>
                        </td>
                      </tr>
                    </table>
                  </td>
                </tr>

                <tr>
                  <td style="padding:14px 18px;border-bottom:1px solid #f0ede5;width:40%;background-color:#faf8f3;">
                    <p style="margin:0;font-size:11px;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;color:{MUTED};">Courier Partner</p>
                  </td>
                  <td style="padding:14px 18px;border-bottom:1px solid #f0ede5;background-color:#ffffff;">
                    <p style="margin:0;font-size:14px;font-weight:700;color:{DARK_GREEN};">{courier_name}</p>
                  </td>
                </tr>

                <tr>
                  <td style="padding:14px 18px;border-bottom:1px solid #f0ede5;background-color:#faf8f3;">
                    <p style="margin:0;font-size:11px;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;color:{MUTED};">Tracking / AWB No.</p>
                  </td>
                  <td style="padding:14px 18px;border-bottom:1px solid #f0ede5;background-color:#ffffff;">
                    <p style="margin:0;font-size:17px;font-weight:900;color:{DARK_GREEN};letter-spacing:0.08em;font-family:Consolas,Monaco,monospace;">
                      {tracking_num}
                    </p>
                  </td>
                </tr>

                <tr>
                  <td style="padding:14px 18px;border-bottom:1px solid #f0ede5;background-color:#faf8f3;">
                    <p style="margin:0;font-size:11px;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;color:{MUTED};">Bib Number</p>
                  </td>
                  <td style="padding:14px 18px;border-bottom:1px solid #f0ede5;background-color:#ffffff;">
                    <p style="margin:0;font-size:14px;font-weight:800;color:{DARK_GREEN};">{payload['bibNumber']}</p>
                  </td>
                </tr>

                <tr>
                  <td style="padding:14px 18px;background-color:#faf8f3;">
                    <p style="margin:0;font-size:11px;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;color:{MUTED};">Category / Distance</p>
                  </td>
                  <td style="padding:14px 18px;background-color:#ffffff;">
                    <p style="margin:0;font-size:14px;font-weight:700;color:{DARK_GREEN};">{payload.get('distance', 'Sports Day Finisher')}</p>
                  </td>
                </tr>
              </table>
            </td>
          </tr>

          <!-- CTA Button -->
          <tr>
            <td style="padding:0 36px 28px;text-align:center;">
              <table width="100%" cellpadding="0" cellspacing="0" border="0">
                <tr>
                  <td align="center">
                    <a href="{tracking_url}" target="_blank" style="display:inline-block;background-color:{DARK_GREEN};color:#ffffff;font-size:14px;font-weight:800;letter-spacing:0.08em;text-transform:uppercase;text-decoration:none;padding:15px 36px;border-radius:10px;box-shadow:0 6px 18px rgba(26,58,46,0.25);border:1px solid {GOLD};">
                      🚚 Track Your Medal Live &rarr;
                    </a>
                  </td>
                </tr>
                <tr>
                  <td align="center" style="padding-top:10px;">
                    <p style="margin:0;font-size:11px;color:#777777;">
                      Direct Link: <a href="{tracking_url}" target="_blank" style="color:{DARK_GREEN};font-weight:600;word-break:break-all;">{tracking_url}</a>
                    </p>
                  </td>
                </tr>
              </table>
            </td>
          </tr>

          {f'''<!-- Delivery Address -->
          <tr>
            <td style="padding:0 36px 24px;">
              <table width="100%" cellpadding="0" cellspacing="0" border="0" style="background-color:#faf8f3;border:1px dashed {LINE};border-radius:10px;padding:14px 18px;">
                <tr>
                  <td>
                    <p style="margin:0 0 4px;font-size:11px;font-weight:800;letter-spacing:0.1em;text-transform:uppercase;color:{MUTED};">📍 Shipping Address Provided</p>
                    <p style="margin:0;font-size:13px;color:#444444;line-height:1.5;">{address_display}</p>
                  </td>
                </tr>
              </table>
            </td>
          </tr>''' if address_display else ''}

          <!-- Delivery Guidelines -->
          <tr>
            <td style="padding:0 36px 28px;">
              <table width="100%" cellpadding="0" cellspacing="0" border="0" style="background-color:#ffffff;border:1px solid #ebe5d8;border-radius:10px;padding:18px 20px;">
                <tr>
                  <td>
                    <p style="margin:0 0 10px;font-size:12px;font-weight:800;letter-spacing:0.05em;text-transform:uppercase;color:{DARK_GREEN};">
                      💡 Important Delivery Notes:
                    </p>
                    <ul style="margin:0;padding-left:18px;font-size:13px;color:#555555;line-height:1.7;">
                      <li><strong>Delivery Timeline:</strong> Typically arrives in <strong>4 to 7 working days</strong> depending on your location.</li>
                      <li><strong>Tracking Activation:</strong> Courier tracking status usually reflects active updates within 24 hours of dispatch.</li>
                      <li><strong>Courier Contact:</strong> Please ensure your registered contact number is reachable for the delivery agent.</li>
                    </ul>
                  </td>
                </tr>
              </table>
            </td>
          </tr>

          <!-- Footer -->
          <tr>
            <td style="background-color:{DARK_GREEN};border-radius:0 0 16px 16px;padding:0;">
              <div style="height:2px;background:linear-gradient(90deg,transparent,{GOLD},transparent);"></div>
              <table width="100%" cellpadding="0" cellspacing="0" border="0">
                <tr>
                  <td style="padding:26px 32px;text-align:center;">
                    <p style="margin:0 0 4px;font-size:17px;color:rgba(255,255,255,0.95);font-family:Georgia,serif;font-style:italic;">
                      Keep Running, Keep Inspiring!
                    </p>
                    <p style="margin:0 0 16px;font-size:10px;color:{GOLD};letter-spacing:0.2em;text-transform:uppercase;">
                      ── Every Finish Has a Story ──
                    </p>
                    <p style="margin:0 0 2px;font-size:13px;font-weight:700;color:#ffffff;">Mountain Run Team</p>
                    <p style="margin:0 0 16px;font-size:11px;color:rgba(255,255,255,0.5);">
                      Official Virtual Marathon Platform · <a href="https://mountainrun.in" target="_blank" style="color:{GOLD};text-decoration:none;">mountainrun.in</a>
                    </p>
                    <table width="180" cellpadding="0" cellspacing="0" border="0" align="center" style="margin:0 auto;">
                      <tr>
                        <td width="60" height="3" style="background-color:#FF9933;border-radius:2px 0 0 2px;"></td>
                        <td width="60" height="3" style="background-color:#ffffff;"></td>
                        <td width="60" height="3" style="background-color:#138808;border-radius:0 2px 2px 0;"></td>
                      </tr>
                    </table>
                  </td>
                </tr>
              </table>
            </td>
          </tr>

          <!-- Disclaimer / Support -->
          <tr>
            <td style="padding:16px;text-align:center;">
              <p style="margin:0;font-size:10px;color:#94a3b8;line-height:1.6;">
                Need help with your shipment? Reach us at <a href="mailto:mountainrunofficial@gmail.com" style="color:#64748b;text-decoration:underline;">mountainrunofficial@gmail.com</a>.<br/>
                © {datetime.datetime.now().year} Mountain Run. All rights reserved.
              </p>
            </td>
          </tr>

        </table>
      </td>
    </tr>
  </table>
</body>
</html>
"""

def send_resend_email(to_email, subject, html_content, max_retries=3):
    if not RESEND_API_KEY:
        raise ValueError("RESEND_API_KEY not found in backend/.env")
    url = "https://api.resend.com/emails"
    headers = {
        "Authorization": f"Bearer {RESEND_API_KEY}",
        "Content-Type": "application/json",
        "User-Agent": "MountainRunDispatch/1.0",
    }
    payload = {
        "from": RESEND_FROM,
        "to": [to_email],
        "subject": subject,
        "html": html_content,
        "reply_to": "mountainrunofficial@gmail.com",
    }
    for attempt in range(max_retries):
        try:
            req = urllib.request.Request(url, data=json.dumps(payload).encode("utf-8"), headers=headers, method="POST")
            with urllib.request.urlopen(req) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            if e.code == 429 and attempt < max_retries - 1:
                wait_s = (attempt + 1) * 2
                print(f"       ⏳ [429 Rate-Limit] Backing off for {wait_s}s...")
                time.sleep(wait_s)
                continue
            err_msg = e.read().decode("utf-8") if e.fp else str(e)
            raise RuntimeError(f"HTTP {e.code}: {err_msg}")

def clean_val(v):
    if v is None: return ""
    s = str(v).strip()
    return "" if s in ["—", "-", "--", "None", "null", "", "?", ""] else s

def run_broadcast():
    print("=================================================================")
    print("   MOUNTAIN RUN — SPORTS DAY MEDAL DISPATCH EMAIL BROADCAST      ")
    print("=================================================================\n")

    excel_path = "tracking id list.xlsx"
    if not os.path.exists(excel_path):
        print(f"Error: {excel_path} not found!")
        return

    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Dispatch List"]

    # 1. Collect all rows that have tracking ID
    runners = []
    for r in range(4, ws.max_row + 1):
        bib = clean_val(ws.cell(row=r, column=2).value)
        awb = clean_val(ws.cell(row=r, column=18).value)
        if bib and awb:
            name = clean_val(ws.cell(row=r, column=4).value) or clean_val(ws.cell(row=r, column=3).value)
            email = clean_val(ws.cell(row=r, column=6).value)
            dist = clean_val(ws.cell(row=r, column=7).value)
            courier = clean_val(ws.cell(row=r, column=19).value)
            link = clean_val(ws.cell(row=r, column=20).value)
            line1 = clean_val(ws.cell(row=r, column=10).value)
            line2 = clean_val(ws.cell(row=r, column=11).value)
            city = clean_val(ws.cell(row=r, column=12).value)
            state = clean_val(ws.cell(row=r, column=13).value)
            pincode = clean_val(ws.cell(row=r, column=14).value)

            runners.append({
                "row": r,
                "bib": bib,
                "name": name,
                "email": email,
                "distance": dist,
                "courier": courier or "Courier Partner",
                "awb": awb,
                "link": link or "https://mountainrun.in",
                "line1": line1,
                "line2": line2,
                "city": city,
                "state": state,
                "pincode": pincode
            })

    print(f"Found {len(runners)} athletes with valid tracking numbers in '{excel_path}'.\n")

    # Connect to DB to update MedalDelivery status
    conn = None
    for attempt in range(5):
        try:
            conn = mm.get_db_connection()
            break
        except Exception as e:
            print(f"DB connection attempt {attempt+1}/5 failed: {e}")
            time.sleep(2)
    cur = conn.cursor(cursor_factory=RealDictCursor) if conn else None

    results = []
    sent_count = 0
    failed_count = 0

    for idx, r in enumerate(runners, start=1):
        payload = {
            "runnerName": r["name"],
            "bibNumber": r["bib"],
            "eventTitle": "Sports Day Celebration 2026",
            "distance": f"{r['distance']} Challenge" if r["distance"] else "Sports Day Finisher",
            "courier": r["courier"],
            "trackingNumber": r["awb"],
            "trackingUrl": r["link"],
            "shippingLine1": r["line1"],
            "shippingLine2": r["line2"],
            "shippingCity": r["city"],
            "shippingState": r["state"],
            "shippingPincode": r["pincode"],
        }

        subject = f"Your Finisher Medal is on the Way! 🏅 — {payload['eventTitle']}"
        html = build_sportsday_dispatch_html(payload)

        print(f"[{idx:2d}/{len(runners)}] Sending to {r['name']} <{r['email']}> | BIB: {r['bib']} | AWB: {r['awb']} ({r['courier']})...", end=" ")

        try:
            res = send_resend_email(r["email"], subject, html)
            res_id = res.get("id")
            print(f"✓ Sent! (ID: {res_id})")
            sent_count += 1
            results.append({
                "bib": r["bib"],
                "name": r["name"],
                "email": r["email"],
                "awb": r["awb"],
                "courier": r["courier"],
                "resendId": res_id,
                "status": "SENT"
            })

            # Update DB MedalDelivery status if connected
            if cur:
                try:
                    cur.execute("""
                        INSERT INTO "MedalDelivery" ("id", "registrationId", "status", "courier", "trackingNumber", "dispatchedAt")
                        SELECT gen_random_uuid(), r.id, 'DISPATCHED', %s, %s, NOW()
                        FROM "Registration" r
                        WHERE r."bibNumber" = %s
                        ON CONFLICT ("registrationId") DO UPDATE
                        SET "status" = 'DISPATCHED', "courier" = EXCLUDED."courier", "trackingNumber" = EXCLUDED."trackingNumber", "dispatchedAt" = NOW()
                    """, (r["courier"], r["awb"], r["bib"]))
                    conn.commit()
                except Exception as dbe:
                    conn.rollback()
                    # Non-fatal: email was sent successfully
                    pass

        except Exception as ex:
            print(f"❌ FAILED: {ex}")
            failed_count += 1
            results.append({
                "bib": r["bib"],
                "name": r["name"],
                "email": r["email"],
                "awb": r["awb"],
                "courier": r["courier"],
                "error": str(ex),
                "status": "FAILED"
            })

        # Rate-limiting pause (700ms)
        time.sleep(0.7)

    if conn:
        conn.close()

    # Save results to JSON log
    log_file = "sportsday_dispatch_broadcast_results.json"
    with open(log_file, "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2, ensure_ascii=False)

    print("\n=================================================================")
    print(f"   BROADCAST COMPLETE: Total {len(runners)} | Sent: {sent_count} | Failed: {failed_count}")
    print(f"   Detailed logs written to: {log_file}")
    print("=================================================================")

if __name__ == "__main__":
    run_broadcast()
