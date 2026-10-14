import os
import sys
import datetime
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from psycopg2.extras import RealDictCursor

sys.path.insert(0, os.path.abspath("."))
import scripts.medal_dispatch_manager as mm

def format_time(sec):
    if not sec:
        return "—"
    try:
        sec = int(sec)
        h = sec // 3600
        m = (sec % 3600) // 60
        s = sec % 60
        if h > 0:
            return f"{h:02d}:{m:02d}:{s:02d}"
        return f"{m:02d}:{s:02d}"
    except Exception:
        return str(sec)

def format_date(dt):
    if not dt:
        return "—"
    if isinstance(dt, (datetime.datetime, datetime.date)):
        return dt.strftime("%d-%b-%Y %H:%M")
    return str(dt)

def format_short_date(dt):
    if not dt:
        return "—"
    if isinstance(dt, (datetime.datetime, datetime.date)):
        return dt.strftime("%d-%b-%Y")
    return str(dt)

def generate_remaining_approved_list():
    # 1. Read existing BIBs in SPORTS_DAY_DISPATCH_LIST_v2.xlsx
    v2_path = "SPORTS_DAY_DISPATCH_LIST_v2.xlsx"
    v2_bibs = set()
    if os.path.exists(v2_path):
        wb_v2 = openpyxl.load_workbook(v2_path)
        ws_v2 = wb_v2["Dispatch List"]
        for r in range(4, ws_v2.max_row + 1):
            b = ws_v2.cell(row=r, column=2).value
            if b:
                v2_bibs.add(str(b).strip().upper())
    print(f"Loaded {len(v2_bibs)} existing dispatched BIBs from {v2_path}")

    # 2. Query Database with retries
    conn = None
    for attempt in range(5):
        try:
            conn = mm.get_db_connection()
            break
        except Exception as e:
            print(f"DB retry {attempt+1}/5: {e}")
            time.sleep(2)

    if not conn:
        raise RuntimeError("Failed to connect to database after 5 attempts")

    cur = conn.cursor(cursor_factory=RealDictCursor)

    cur.execute("""
        SELECT 
            r.id as registration_id,
            r."bibNumber" as bib_number,
            r."userId" as user_id,
            u.name as user_name,
            u.email as email,
            COALESCE(r."shippingPhone", u.phone) as phone,
            r.distance as distance,
            r."activityType" as activity_type,
            r."finishTimeSeconds" as finish_time_seconds,
            r."shippingName" as shipping_name,
            r."shippingLine1" as address_line1,
            COALESCE(r."shippingLine2", '') as address_line2,
            r."shippingCity" as city,
            r."shippingState" as state,
            r."shippingPincode" as pincode,
            r."proofStatus" as proof_status,
            COALESCE(m.status, 'PENDING') as medal_status,
            m."trackingNumber" as tracking_number,
            m."courier" as courier,
            m."dispatchedAt" as dispatched_at,
            r."registeredAt" as registered_at,
            c."certificateNumber" as certificate_number,
            pu."activityImageUrl" as proof_url,
            pu."submittedAt" as proof_submitted_at,
            pu."reviewedAt" as proof_reviewed_at,
            p."razorpayPaymentId" as payment_id,
            p."amountInPaise" as amount_in_paise,
            p."paidAt" as paid_at
        FROM "Registration" r
        JOIN "User" u ON r."userId" = u.id
        JOIN "Event" e ON r."eventId" = e.id
        LEFT JOIN "Payment" p ON p."registrationId" = r.id
        LEFT JOIN "MedalDelivery" m ON m."registrationId" = r.id
        LEFT JOIN "ProofUpload" pu ON pu."registrationId" = r.id
        LEFT JOIN "Certificate" c ON c."registrationId" = r.id
        WHERE e.slug = 'sports-day-celebration' AND r."proofStatus" = 'APPROVED'
        ORDER BY r."registeredAt" ASC
    """)
    all_approved = cur.fetchall()
    conn.close()

    # Filter for remaining
    remaining = [r for r in all_approved if r["bib_number"].strip().upper() not in v2_bibs]
    print(f"Total Approved in DB: {len(all_approved)} | Remaining to dispatch: {len(remaining)}")

    # 3. Styling definitions matching v2
    DARK_GREEN    = "1A5C38"   # header banner
    MED_GREEN     = "2D7D52"   # sub banner / col headers
    LIGHT_GREEN   = "E8F5EE"   # even rows
    WHITE         = "FFFFFF"   # odd rows
    BORDER_COLOR  = "C8E6C9"   # grid lines
    HEADER_TEXT   = "FFFFFF"
    BODY_TEXT     = "1A1A1A"

    def make_fill(hex_c): return PatternFill(fill_type="solid", fgColor=hex_c)
    def make_border(color=BORDER_COLOR, style="thin"):
        s = Side(style=style, color=color)
        return Border(left=s, right=s, top=s, bottom=s)
    def header_font(size=9, bold=True): return Font(name="Calibri", size=size, bold=bold, color=HEADER_TEXT)
    def body_font(size=9, bold=False, color=BODY_TEXT, italic=False):
        return Font(name="Calibri", size=size, bold=bold, color=color, italic=italic)
    def center_align(wrap=True): return Alignment(horizontal="center", vertical="center", wrap_text=wrap)
    def left_align(wrap=True): return Alignment(horizontal="left", vertical="center", wrap_text=wrap)

    # 4. Create Workbook
    wb = openpyxl.Workbook()

    # ─────────────────────────────────────────────────────────────────
    # SHEET 1: Dispatch List
    # ─────────────────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Dispatch List"
    ws1.freeze_panes = "D4"

    # Banner Row 1
    ws1.merge_cells("A1:W1")
    b1 = ws1["A1"]
    b1.value = "MOUNTAIN RUN — SPORTS DAY CELEBRATION 2026"
    b1.font = Font(name="Calibri", size=15, bold=True, color=HEADER_TEXT)
    b1.fill = make_fill(DARK_GREEN)
    b1.alignment = center_align()
    ws1.row_dimensions[1].height = 32

    # Banner Row 2
    ws1.merge_cells("A2:W2")
    b2 = ws1["A2"]
    b2.value = f"Remaining Approved Medal Dispatch (Batch 2) — {len(remaining)} Athletes  |  Generated: 08-Sep-2026"
    b2.font = Font(name="Calibri", size=10, bold=False, italic=True, color=HEADER_TEXT)
    b2.fill = make_fill(MED_GREEN)
    b2.alignment = center_align()
    ws1.row_dimensions[2].height = 20

    COLS = [
        ("S.No", 6),
        ("BIB Number", 13),
        ("Runner Name", 20),
        ("Shipping Recipient", 20),
        ("Mobile", 14),
        ("Email", 28),
        ("Distance", 10),
        ("Activity", 10),
        ("Finish Time", 12),
        ("Address Line 1", 32),
        ("Address Line 2", 22),
        ("City", 15),
        ("State", 16),
        ("Pincode", 9),
        ("Proof Status", 13),
        ("Proof Reviewed Date", 18),
        ("Medal Status", 14),
        ("AWB / Tracking", 16),
        ("Courier", 13),
        ("Certificate No", 20),
        ("Reg Date", 13),
        ("Amount (Rs)", 12),
        ("Payment ID", 22),
    ]

    for i, (label, w) in enumerate(COLS, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w
        c = ws1.cell(row=3, column=i, value=label)
        c.font = header_font(9)
        c.fill = make_fill("2E7D52")
        c.alignment = center_align()
        c.border = make_border("0D4A2A", "thin")
    ws1.row_dimensions[3].height = 24

    for row_idx, r in enumerate(remaining, start=1):
        excel_row = row_idx + 3
        is_even = (row_idx % 2 == 0)
        row_bg = LIGHT_GREEN if is_even else WHITE

        amt = r["amount_in_paise"] / 100 if r["amount_in_paise"] else 0
        amt_str = f"Rs. {int(amt)}" if amt > 0 else "—"

        row_values = [
            row_idx,
            r["bib_number"],
            r["user_name"],
            r["shipping_name"] or r["user_name"],
            r["phone"] or "—",
            r["email"],
            r["distance"],
            r["activity_type"].capitalize() if r["activity_type"] else "Running",
            format_time(r["finish_time_seconds"]),
            r["address_line1"] or "—",
            r["address_line2"] or "—",
            r["city"] or "—",
            r["state"] or "—",
            r["pincode"] or "—",
            r["proof_status"],
            format_date(r["proof_reviewed_at"] or r["proof_submitted_at"]),
            r["medal_status"] or "PENDING",
            r["tracking_number"] or "—",
            r["courier"] or "—",
            r["certificate_number"] or "—",
            format_short_date(r["registered_at"]),
            amt_str,
            r["payment_id"] or "—",
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws1.cell(row=excel_row, column=col_idx, value=val)
            cell.border = make_border()
            is_center = col_idx in [1, 2, 5, 7, 8, 9, 14, 15, 16, 17, 18, 19, 21, 22]
            cell.alignment = center_align(wrap=True) if is_center else left_align(wrap=True)

            # Highlighting
            if col_idx == 15:  # Proof Status
                cell.fill = make_fill("C8E6C9")
                cell.font = body_font(bold=True, color="1B5E20")
            elif col_idx == 17: # Medal Status
                cell.fill = make_fill("FFF9C4")  # yellow pending
                cell.font = body_font(bold=True, color="F57F17")
            else:
                cell.fill = make_fill(row_bg)
                cell.font = body_font(color="555555" if val == "—" else BODY_TEXT, italic=(val == "—"))

        ws1.row_dimensions[excel_row].height = 22

    ws1.auto_filter.ref = f"A3:W{3 + len(remaining)}"
    ws1.sheet_view.showGridLines = True

    # ─────────────────────────────────────────────────────────────────
    # SHEET 2: Courier Ready
    # ─────────────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Courier Ready")
    ws2.freeze_panes = "A3"

    ws2.merge_cells("A1:J1")
    c1 = ws2["A1"]
    c1.value = "COURIER DISPATCH SHEET — SPORTS DAY 2026 (BATCH 2)"
    c1.font = Font(name="Calibri", size=13, bold=True, color=HEADER_TEXT)
    c1.fill = make_fill(DARK_GREEN)
    c1.alignment = center_align()
    ws2.row_dimensions[1].height = 26

    COURIER_COLS = [
        ("S.No", 6),
        ("BIB No", 13),
        ("Runner Name", 20),
        ("Mobile", 14),
        ("Full Shipping Address", 45),
        ("City", 14),
        ("State", 16),
        ("Pincode", 10),
        ("Medal Status", 14),
        ("AWB / Tracking", 16),
    ]

    for i, (label, w) in enumerate(COURIER_COLS, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
        cell = ws2.cell(row=2, column=i, value=label)
        cell.font = header_font(9)
        cell.fill = make_fill(MED_GREEN)
        cell.alignment = center_align()
        cell.border = make_border("0D4A2A")
    ws2.row_dimensions[2].height = 20

    for row_idx, r in enumerate(remaining, start=1):
        er = row_idx + 2
        is_even = (row_idx % 2 == 0)
        bg = LIGHT_GREEN if is_even else WHITE

        full_addr = f"{r['address_line1'] or ''}, {r['address_line2'] or ''}".strip(", ")

        vals = [
            row_idx,
            r["bib_number"],
            r["shipping_name"] or r["user_name"],
            r["phone"] or "—",
            full_addr,
            r["city"] or "—",
            r["state"] or "—",
            r["pincode"] or "—",
            r["medal_status"] or "PENDING",
            r["tracking_number"] or "—",
        ]

        for ci, v in enumerate(vals, start=1):
            cell = ws2.cell(row=er, column=ci, value=v)
            cell.border = make_border()
            if ci == 9:
                cell.fill = make_fill("FFF9C4")
                cell.font = body_font(bold=True, color="F57F17")
            else:
                cell.fill = make_fill(bg)
                cell.font = body_font()
            cell.alignment = center_align() if ci in [1, 2, 4, 6, 7, 8, 9, 10] else left_align()
        ws2.row_dimensions[er].height = 20

    ws2.auto_filter.ref = f"A2:J{2 + len(remaining)}"
    ws2.sheet_view.showGridLines = True

    # ─────────────────────────────────────────────────────────────────
    # SHEET 3: Summary Dashboard
    # ─────────────────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Summary Dashboard")
    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 20
    ws3.column_dimensions["C"].width = 18
    ws3.column_dimensions["D"].width = 18

    ws3.merge_cells("A1:D1")
    ws3["A1"].value = "SPORTS DAY — BATCH 2 DISPATCH DASHBOARD"
    ws3["A1"].font = Font(name="Calibri", size=14, bold=True, color=HEADER_TEXT)
    ws3["A1"].fill = make_fill(DARK_GREEN)
    ws3["A1"].alignment = center_align()
    ws3.row_dimensions[1].height = 28

    def dash_section(ws, row, title):
        ws.merge_cells(f"A{row}:D{row}")
        c = ws.cell(row=row, column=1, value=title)
        c.font = Font(name="Calibri", size=10, bold=True, color=HEADER_TEXT)
        c.fill = make_fill(MED_GREEN)
        c.alignment = center_align()
        c.border = make_border(MED_GREEN)
        ws.row_dimensions[row].height = 18

    def dash_row(ws, row, label, value, bg=WHITE):
        lc = ws.cell(row=row, column=1, value=label)
        vc = ws.cell(row=row, column=2, value=value)
        lc.font = body_font(bold=True)
        vc.font = body_font(color="1B5E20", bold=True)
        lc.fill = vc.fill = make_fill(bg)
        lc.alignment = left_align(False)
        vc.alignment = center_align(False)
        lc.border = vc.border = make_border()
        ws.row_dimensions[row].height = 18
        for c in [3, 4]:
            xc = ws.cell(row=row, column=c)
            xc.fill = make_fill(bg)
            xc.border = make_border()

    dist_counts = {}
    for r in remaining:
        d = r["distance"]
        dist_counts[d] = dist_counts.get(d, 0) + 1

    r_idx = 2
    dash_section(ws3, r_idx, "BATCH 2 OVERVIEW"); r_idx += 1
    dash_row(ws3, r_idx, "Remaining Approved Athletes", len(remaining), LIGHT_GREEN); r_idx += 1
    dash_row(ws3, r_idx, "Total Previously Dispatched (Batch 1)", len(v2_bibs), WHITE); r_idx += 1
    dash_row(ws3, r_idx, "Total Approved Overall in DB", len(all_approved), LIGHT_GREEN); r_idx += 1

    r_idx += 1
    dash_section(ws3, r_idx, "BREAKDOWN BY DISTANCE (BATCH 2)"); r_idx += 1
    for dist, cnt in sorted(dist_counts.items()):
        dash_row(ws3, r_idx, dist, cnt, LIGHT_GREEN if r_idx % 2 == 0 else WHITE); r_idx += 1

    r_idx += 1
    dash_section(ws3, r_idx, "MEDAL STATUS"); r_idx += 1
    dash_row(ws3, r_idx, "Pending Dispatch", len(remaining), "FFF9C4"); r_idx += 1

    ws3.sheet_view.showGridLines = True

    # 5. Save Excel File
    output_excel = "SPORTS_DAY_REMAINING_APPROVED_DISPATCH_LIST.xlsx"
    wb.save(output_excel)
    print(f"\n[DONE] Saved new Excel file to: {output_excel}")

    # Also export to PDF
    try:
        from scripts.xlsx_to_pdf import xlsx_to_pdf
        output_pdf = "SPORTS_DAY_REMAINING_APPROVED_DISPATCH_LIST.pdf"
        # Temporarily create clean 1-sheet for PDF
        temp_wb = openpyxl.Workbook()
        temp_ws = temp_wb.active
        temp_ws.title = "Dispatch"
        # Copy header and rows from Courier Ready
        for row in ws2.iter_rows(values_only=True):
            temp_ws.append(list(row))
        temp_excel = "temp_batch2_for_pdf.xlsx"
        temp_wb.save(temp_excel)
        xlsx_to_pdf(temp_excel, output_pdf)
        if os.path.exists(temp_excel): os.remove(temp_excel)
        print(f"[DONE] Saved corresponding PDF to: {output_pdf}")
    except Exception as ex:
        print(f"[PDF Note] Could not auto-generate PDF: {ex}")

if __name__ == "__main__":
    generate_remaining_approved_list()
