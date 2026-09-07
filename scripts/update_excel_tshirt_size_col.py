import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

wb = openpyxl.load_workbook('SPORTS_DAY_LEADERBOARD_TSHIRT_WINNERS.xlsx')
ws = wb['Top 3 T-Shirt Winners']

# Add data validation dropdown for T-Shirt Size in Column M (rows 5 to 25)
dv = DataValidation(
    type="list",
    formula1='"S, M, L, XL, XXL"',
    allow_blank=True,
    showInputMessage=True,
    showErrorMessage=True
)
dv.errorTitle = 'Invalid T-Shirt Size'
dv.error = 'Please select a valid size: S, M, L, XL, or XXL'
dv.promptTitle = 'Select T-Shirt Size'
dv.prompt = 'Choose runner size: S, M, L, XL, or XXL'

ws.add_data_validation(dv)
dv.add("M5:M25")

# Format Column M (T-Shirt Size) with a clear, subtle highlight so it's easy to spot for input
input_fill = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid") # subtle warm cream/amber
bold_font = Font(name="Calibri", size=10, bold=True, color="000000")
center_align = Alignment(horizontal="center", vertical="center")

thin_border = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB')
)

for row in range(5, 26):
    cell = ws.cell(row=row, column=13)
    cell.fill = input_fill
    cell.font = bold_font
    cell.alignment = center_align
    cell.border = thin_border

# Save workbook
wb.save('SPORTS_DAY_LEADERBOARD_TSHIRT_WINNERS.xlsx')
print("Successfully enhanced SPORTS_DAY_LEADERBOARD_TSHIRT_WINNERS.xlsx with dropdown data validation & clear styling!")
