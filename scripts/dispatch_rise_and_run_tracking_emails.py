import os
import sys
import json
import time
import datetime
import urllib.request
import urllib.error
import openpyxl
import psycopg2
from psycopg2.extras import RealDictCursor

# Add repository root to path
sys.path.insert(0, os.path.abspath("."))
import scripts.medal_dispatch_manager as mm

ENV = mm.get_backend_env()
RESEND_API_KEY = ENV.get("RESEND_API_KEY")
RESEND_FROM = ENV.get("RESEND_FROM_EMAIL", "Mountain Run <onboarding@mountainrun.in>")
DB_URL = ENV.get("DATABASE_URL")

def get_db_connection(max_attempts=5):
    for attempt in range(max_attempts):
        try:
            return psycopg2.connect(DB_URL, connect_timeout=15)
        except Exception as e:
            if attempt < max_attempts - 1:
                time.sleep(2)
            else:
                raise e

def get_courier_details(raw_courier, tracking_number):
    c_upper = str(raw_courier or "").upper().strip()
    t_str = str(tracking_number or "").strip()

    if "DALEVER" in c_upper or "DELHIVERY" in c_upper:
        return {
            "name": "Delhivery Express",
            "url": f"https://www.delhivery.com/track/package/{t_str}"
        }
    elif "BLUEDART" in c_upper or "BLUE DART" in c_upper:
        return {
            "name": "Blue Dart Express",
            "url": "https://www.bluedart.com/tracking"
        }
    else:
        return {
            "name": "DTDC Express",
            "url": "https://www.dtdc.com/track-your-shipment/"
        }

def send_resend_email(to_email, subject, html_content, max_retries=3):
    if not RESEND_API_KEY:
        raise ValueError("RESEND_API_KEY not found in backend/.env")
    url = "https://api.resend.com/emails"
    headers = {
        "Authorization": f"Bearer {RESEND_API_KEY}",
        "Content-Type": "application/json",
        "User-Agent": "MountainRunDispatch/1.0",
    }
    payload = {
        "from": RESEND_FROM,
        "to": [to_email],
        "subject": subject,
        "html": html_content,
        "reply_to": "mountainrunofficial@gmail.com",
    }
    for attempt in range(max_retries):
        try:
            req = urllib.request.Request(url, data=json.dumps(payload).encode("utf-8"), headers=headers, method="POST")
            with urllib.request.urlopen(req) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            if e.code == 429 and attempt < max_retries - 1:
                wait_s = (attempt + 1) * 2
                print(f"       ⏳ [429 Rate-Limit] Waiting {wait_s}s...")
                time.sleep(wait_s)
                continue
            err_msg = e.read().decode("utf-8") if e.fp else str(e)
            raise RuntimeError(f"HTTP {e.code}: {err_msg}")

def run_dispatch(live=False):
    excel_path = os.path.join(os.getcwd(), "SAINI.xlsx")
    if not os.path.exists(excel_path):
        print(f"❌ Error: {excel_path} not found!")
        return

    print("=" * 65)
    mode_str = "LIVE BROADCAST & DB SYNC" if live else "DRY-RUN SIMULATION (NO EMAILS SENT)"
    print(f"   MOUNTAIN RUN — RISE AND RUN MEDAL DISPATCH [{mode_str}]")
    print("=" * 65 + "\n")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb["Courier Dispatch Ready"]

    # 1. Parse Excel rows
    candidates = []
    for r in range(4, sheet.max_row + 1):
        sno = sheet.cell(r, 1).value
        bib = str(sheet.cell(r, 2).value or "").strip()
        name = str(sheet.cell(r, 3).value or "").strip()
        phone = str(sheet.cell(r, 4).value or "").strip()
        email = str(sheet.cell(r, 6).value or "").strip()
        tracking = str(sheet.cell(r, 16).value or "").strip()
        courier_raw = str(sheet.cell(r, 18).value or "").strip()

        if not bib and not name:
            continue

        if not tracking or tracking.lower() in ["none", "nan", "null", "-", ""]:
            continue

        candidates.append({
            "row": r,
            "sno": sno,
            "bib": bib,
            "sheet_name": name,
            "sheet_phone": phone,
            "sheet_email": email,
            "tracking": tracking,
            "courier_raw": courier_raw
        })

    print(f"📦 Total athletes with tracking IDs parsed from SAINI.xlsx: {len(candidates)}")

    # 2. Connect to DB and fetch participant data
    conn = get_db_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)

    cur.execute("""
        SELECT 
            r.id as registration_id,
            r."bibNumber" as bib_number,
            r.distance,
            r."shippingName" as shipping_name,
            r."shippingPhone" as shipping_phone,
            r."shippingLine1" as address_line1,
            COALESCE(r."shippingLine2", '') as address_line2,
            r."shippingCity" as city,
            r."shippingState" as state,
            r."shippingPincode" as pincode,
            u.name as user_name,
            u.email as user_email,
            e.title as event_title,
            m.id as medal_delivery_id,
            m.status as medal_status,
            m."trackingNumber" as db_tracking_number,
            m.courier as db_courier
        FROM "Registration" r
        JOIN "User" u ON r."userId" = u.id
        JOIN "Event" e ON r."eventId" = e.id
        LEFT JOIN "MedalDelivery" m ON m."registrationId" = r.id
        WHERE e.slug ILIKE '%rise%' OR r."bibNumber" ILIKE 'RAR-%'
    """)
    db_rows = cur.fetchall()
    db_by_bib = {r["bib_number"].strip().upper(): r for r in db_rows if r["bib_number"]}

    print(f"🗄️ Fetched {len(db_rows)} Rise and Run registrations from DB.\n")

    logs = []
    success_count = 0
    skip_count = 0
    fail_count = 0

    now_iso = datetime.datetime.now(datetime.timezone.utc)

    for i, c in enumerate(candidates, 1):
        bib_upper = c["bib"].upper()
        db_reg = db_by_bib.get(bib_upper)

        if not db_reg:
            print(f"[{i:02d}/{len(candidates)}] ❌ BIB {c['bib']} ({c['sheet_name']}) not found in DB! Skipping.")
            fail_count += 1
            logs.append({"bib": c["bib"], "status": "NOT_FOUND_IN_DB"})
            continue

        runner_name = db_reg["shipping_name"] or db_reg["user_name"] or c["sheet_name"]
        to_email = db_reg["user_email"] or c["sheet_email"]
        courier_info = get_courier_details(c["courier_raw"], c["tracking"])
        dist_str = f"{db_reg['distance']} Challenge" if db_reg['distance'] else "Finisher"

        payload = {
            "to": to_email,
            "runnerName": runner_name,
            "bibNumber": c["bib"],
            "eventTitle": db_reg["event_title"] or "Rise and Run Virtual Challenge 2026",
            "distance": dist_str,
            "courier": courier_info["name"],
            "trackingNumber": c["tracking"],
            "trackingUrl": courier_info["url"],
            "shippingLine1": db_reg["address_line1"],
            "shippingLine2": db_reg["address_line2"],
            "shippingCity": db_reg["city"],
            "shippingState": db_reg["state"],
            "shippingPincode": db_reg["pincode"],
        }

        # Check if already dispatched with same tracking
        if db_reg.get("medal_status") == "DISPATCHED" and db_reg.get("db_tracking_number") == c["tracking"]:
            print(f"[{i:02d}/{len(candidates)}] ⏩ Already dispatched with this tracking: {c['bib']} | {runner_name}")
            skip_count += 1
            logs.append({"bib": c["bib"], "status": "ALREADY_DISPATCHED", "email": to_email})
            continue

        if not live:
            # DRY RUN
            print(f"[{i:02d}/{len(candidates)}] [DRY RUN] Would send to: {runner_name} <{to_email}> | {courier_info['name']} | AWB: {c['tracking']}")
            success_count += 1
            logs.append({
                "bib": c["bib"],
                "runner": runner_name,
                "email": to_email,
                "courier": courier_info["name"],
                "tracking": c["tracking"],
                "status": "DRY_RUN_OK"
            })
        else:
            # LIVE EXECUTION
            subject = f"Your Finisher Medal is on the Way! 🏅 — {payload['eventTitle']}"
            html = mm.build_medal_dispatch_html(payload)

            print(f"[{i:02d}/{len(candidates)}] Sending to {runner_name} <{to_email}> [{courier_info['name']} {c['tracking']}]...", end=" ", flush=True)

            try:
                # 1. Send Email
                res = send_resend_email(to_email, subject, html)
                res_id = res.get("id", "OK")

                # 2. Update DB
                cur.execute("""
                    INSERT INTO "MedalDelivery" ("id", "registrationId", "status", "courier", "trackingNumber", "trackingUrl", "dispatchedAt")
                    VALUES (gen_random_uuid(), %s, 'DISPATCHED', %s, %s, %s, %s)
                    ON CONFLICT ("registrationId") DO UPDATE
                    SET "status" = 'DISPATCHED',
                        "courier" = EXCLUDED."courier",
                        "trackingNumber" = EXCLUDED."trackingNumber",
                        "trackingUrl" = EXCLUDED."trackingUrl",
                        "dispatchedAt" = EXCLUDED."dispatchedAt"
                """, (db_reg["registration_id"], courier_info["name"], c["tracking"], courier_info["url"], now_iso))
                conn.commit()

                print(f"✅ SENT! (Resend ID: {res_id})", flush=True)
                success_count += 1
                logs.append({
                    "bib": c["bib"],
                    "runner": runner_name,
                    "email": to_email,
                    "courier": courier_info["name"],
                    "tracking": c["tracking"],
                    "status": "SENT",
                    "resend_id": res_id
                })

                # Slight throttle to prevent Resend 429
                time.sleep(0.35)

            except Exception as ex:
                conn.rollback()
                print(f"❌ FAILED: {ex}", flush=True)
                fail_count += 1
                logs.append({
                    "bib": c["bib"],
                    "runner": runner_name,
                    "email": to_email,
                    "courier": courier_info["name"],
                    "tracking": c["tracking"],
                    "status": "FAILED",
                    "error": str(ex)
                })

    conn.close()

    # Save results log
    output_log_file = "rise_and_run_dispatch_results.json"
    with open(output_log_file, "w", encoding="utf-8") as f:
        json.dump({
            "timestamp": now_iso.isoformat(),
            "mode": "LIVE" if live else "DRY_RUN",
            "total_candidates": len(candidates),
            "success": success_count,
            "skipped": skip_count,
            "failed": fail_count,
            "details": logs
        }, f, indent=2)

    print("\n" + "=" * 65)
    print(f"   DISPATCH SUMMARY ({mode_str})")
    print(f"   Total Processed : {len(candidates)}")
    print(f"   Success         : {success_count}")
    print(f"   Skipped         : {skip_count}")
    print(f"   Failed          : {fail_count}")
    print(f"   Log saved to    : {output_log_file}")
    print("=" * 65)

if __name__ == "__main__":
    is_live = "--live" in sys.argv
    run_dispatch(live=is_live)
