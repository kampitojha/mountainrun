#!/usr/bin/env python3
"""
Generate RISE_AND_RUN_LEADERBOARD_TSHIRT_WINNERS.xlsx
Creates a premium formatted Excel workbook for tracking Top 3 T-Shirt winners
and live sizes submitted via mountainrun.in/claim-jersey
"""

import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.abspath("scripts"))
from rise_and_run_tshirt_dispatch import get_rise_and_run_top3_winners

winners = get_rise_and_run_top3_winners()

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Top 3 T-Shirt Winners"

# Styles
font_title = Font(name="Calibri", size=16, bold=True, color="1B4D3E")
font_sub = Font(name="Calibri", size=11, italic=True, color="555555")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid")

font_gold = Font(name="Calibri", size=11, bold=True, color="92400E")
fill_gold = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

font_silver = Font(name="Calibri", size=11, bold=True, color="374151")
fill_silver = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

font_bronze = Font(name="Calibri", size=11, bold=True, color="7C2D12")
fill_bronze = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")

font_data = Font(name="Calibri", size=11)
align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")

thin_border = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB')
)

# Title Block
ws.merge_cells("A1:N1")
ws["A1"] = "MOUNTAIN RUN — RISE AND RUN VIRTUAL MARATHON 2026"
ws["A1"].font = font_title
ws["A1"].alignment = align_center

ws.merge_cells("A2:N2")
ws["A2"] = "Official Top 3 Category Winners — DRI-FIT T-Shirt Dispatch List (15 Athletes)"
ws["A2"].font = font_sub
ws["A2"].alignment = align_center

ws.row_dimensions[1].height = 30
ws.row_dimensions[2].height = 20
ws.row_dimensions[4].height = 26

# Headers
headers = [
    "S.No",
    "Category",
    "Rank",
    "BIB Number",
    "Runner Name",
    "Finish Time",
    "Email",
    "Phone",
    "Shipping Address",
    "City",
    "State",
    "Pincode",
    "T-Shirt Size",
    "Claim Status"
]

for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = thin_border

for idx, w in enumerate(winners, 1):
    row_num = idx + 4
    ws.row_dimensions[row_num].height = 22

    status_str = f"Confirmed ({w['tshirt_size']})" if w['tshirt_size'] else "Pending Claim"

    row_data = [
        idx,
        w['cat_raw'],
        w['rank'],
        w['bib'],
        w['name'],
        w['finish_time'],
        w['email'],
        w['phone'],
        w['shipping_address'],
        w.get('city', ''),
        w.get('state', ''),
        w.get('pincode', ''),
        w['tshirt_size'] or "",
        status_str
    ]

    for col_idx, val in enumerate(row_data, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=val)
        cell.font = font_data
        cell.border = thin_border
        
        if col_idx in [1, 2, 4, 6, 8, 12, 13, 14]:
            cell.alignment = align_center
        else:
            cell.alignment = align_left

        if col_idx == 3:  # Rank column formatting
            if "1st" in str(val):
                cell.font = font_gold
                cell.fill = fill_gold
            elif "2nd" in str(val):
                cell.font = font_silver
                cell.fill = fill_silver
            elif "3rd" in str(val):
                cell.font = font_bronze
                cell.fill = fill_bronze

# Auto column width
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.row > 2 and cell.value:
            max_len = max(max_len, len(str(cell.value)))
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

out_file = "RISE_AND_RUN_LEADERBOARD_TSHIRT_WINNERS.xlsx"
wb.save(out_file)
print(f"[✓] Created Excel sheet: {out_file}")
