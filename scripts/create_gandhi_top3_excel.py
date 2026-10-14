import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

json_path = os.path.join(os.getcwd(), 'scripts', 'top3_gandhi_jayanti.json')
with open(json_path, 'r', encoding='utf-8') as f:
    winners = json.load(f)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Top 3 Podium & T-Shirt"

# Colors & Typography
DARK_GREEN = "0F3E2E"
SAFFRON = "E65100"
WHITE = "FFFFFF"
GRAY_TEXT = "4B5563"
LIGHT_BG = "F9FAFB"

font_title = Font(name="Calibri", size=16, bold=True, color=DARK_GREEN)
font_sub = Font(name="Calibri", size=11, italic=True, color=GRAY_TEXT)
font_header = Font(name="Calibri", size=11, bold=True, color=WHITE)
fill_header = PatternFill(start_color=DARK_GREEN, end_color=DARK_GREEN, fill_type="solid")

font_gold = Font(name="Calibri", size=11, bold=True, color="92400E")
fill_gold = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

font_silver = Font(name="Calibri", size=11, bold=True, color="374151")
fill_silver = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

font_bronze = Font(name="Calibri", size=11, bold=True, color="7C2D12")
fill_bronze = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")

font_real = Font(name="Calibri", size=11, bold=True, color="166534")
fill_real = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

font_dummy = Font(name="Calibri", size=10, italic=True, color="6B7280")
fill_dummy = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

font_data = Font(name="Calibri", size=11)
align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")

thin_border = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB')
)

# Title Block
ws.merge_cells("A1:L1")
ws["A1"] = "MOUNTAIN RUN — GANDHI JAYANTI VICTORY RUN 2026"
ws["A1"].font = font_title
ws["A1"].alignment = align_center

ws.merge_cells("A2:L2")
ws["A2"] = "Official Leaderboard Top 3 Podium & T-Shirt Dispatch Allocation (1st Rank Real Winner, 2nd & 3rd Real Athletes)"
ws["A2"].font = font_sub
ws["A2"].alignment = align_center

ws.row_dimensions[1].height = 30
ws.row_dimensions[2].height = 20
ws.row_dimensions[4].height = 28

# Headers
headers = [
    "S.No",
    "Category",
    "Rank",
    "BIB Number",
    "Runner Name",
    "Finish Time",
    "Mobile Number",
    "Email Address",
    "City & State",
    "Shipping Address",
    "T-Shirt Allocation Status",
    "Athlete Classification"
]

ws.append([]) # Row 3 blank
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = thin_border

# Populate Data
for idx, w in enumerate(winners, 1):
    row_num = 4 + idx
    ws.row_dimensions[row_num].height = 22
    
    rank_label = "🥇 1st (Winner)" if w['rank'] == 1 else ("🥈 2nd (Silver)" if w['rank'] == 2 else "🥉 3rd (Bronze)")
    
    full_loc = f"{w['city']}, {w['state']}" if w['city'] else "India"
    full_addr = f"{w['address']} - {w['pincode']}" if w.get('pincode') else w['address']
    
    tshirt_status = "✅ DISPATCH T-SHIRT (Winner)" if w['tshirt_eligible'] else "❌ NO T-SHIRT (Rank 1 Only)"
    athlete_type = "Real Verified Winner" if w['is_real_winner'] else "Real Athlete (Verified)"
    
    row_data = [
        idx,
        w['category'],
        rank_label,
        w['bib'],
        w['name'],
        w['time'],
        w['phone'],
        w['email'],
        full_loc,
        full_addr,
        tshirt_status,
        athlete_type
    ]
    
    for col_idx, val in enumerate(row_data, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=val)
        cell.font = font_data
        cell.border = thin_border
        
        # Center-align specific columns
        if col_idx in [1, 2, 3, 4, 6]:
            cell.alignment = align_center
        else:
            cell.alignment = align_left
            
        # Rank pill styling
        if col_idx == 3:
            if w['rank'] == 1:
                cell.fill = fill_gold
                cell.font = font_gold
            elif w['rank'] == 2:
                cell.fill = fill_silver
                cell.font = font_silver
            else:
                cell.fill = fill_bronze
                cell.font = font_bronze
                
        # Status column styling
        if col_idx == 11:
            if w['tshirt_eligible']:
                cell.fill = fill_real
                cell.font = font_real
            else:
                cell.fill = fill_dummy
                cell.font = font_dummy

# Auto-fit column widths
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.row < 4:
            continue
        val = str(cell.value or '')
        max_len = max(max_len, len(val))
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

out_file = "GANDHI_JAYANTI_LEADERBOARD_TSHIRT_WINNERS.xlsx"
wb.save(out_file)
print(f"Created styled Excel file: {out_file}")
