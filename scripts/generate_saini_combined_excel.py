import os
import sys
import json
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import psycopg2
from psycopg2.extras import RealDictCursor

sys.path.insert(0, os.path.abspath("."))
import scripts.medal_dispatch_manager as mm

def build_combined_saini_excel():
    print("=" * 70)
    print(" GENERATING COMBINED MASTER SHEET: SAINI 2 + SAINI 4")
    print("=" * 70)

    # 1. Load Batch 1 (88)
    with open("rise_and_run_dispatch_results.json", "r", encoding="utf-8") as f:
        b1_raw = json.load(f)["details"]
    b1_map = {x["bib"].strip().upper(): x for x in b1_raw}

    # 2. Load Batch 2 (27)
    with open("rise_and_run_dispatch_batch2_results.json", "r", encoding="utf-8") as f:
        b2_raw = json.load(f)["results"]
    b2_map = {x["bib"].strip().upper(): x for x in b2_raw}

    batch2_first_6 = [
        {"bib": "RAR-744633", "tracking": "7D140782228", "courier": "DTDC Express"},
        {"bib": "RAR-456185", "tracking": "7D140782229", "courier": "DTDC Express"},
        {"bib": "RAR-752636", "tracking": "7D140782238", "courier": "DTDC Express"},
        {"bib": "RAR-104799", "tracking": "42833510728711", "courier": "Delhivery Express"},
        {"bib": "RAR-821656", "tracking": "42833510728814", "courier": "Delhivery Express"},
        {"bib": "RAR-146683", "tracking": "7D140782242", "courier": "DTDC Express"}
    ]
    for item in batch2_first_6:
        b2_map[item["bib"].upper()] = item

    # 3. Load SAINI 4 (19)
    saini4_data = [
        {"bib": "RAR-109452", "awb": "7D140782470", "courier": "DTDC Express", "amount": 95, "status": "DISPATCHED"},
        {"bib": "RAR-815836", "awb": "7D140782475", "courier": "DTDC Express", "amount": 95, "status": "DISPATCHED"},
        {"bib": "RAR-154573", "awb": "7D140782478", "courier": "DTDC Express", "amount": 78, "status": "DISPATCHED"},
        {"bib": "RAR-491655", "awb": "7D140782479", "courier": "DTDC Express", "amount": 78, "status": "DISPATCHED"},
        {"bib": "RAR-107978", "awb": "7D140782480", "courier": "DTDC Express", "amount": 68, "status": "DISPATCHED"},
        {"bib": "RAR-139486", "awb": "7D140782482", "courier": "DTDC Express", "amount": 100, "status": "DISPATCHED"},
        {"bib": "RAR-202469", "awb": "7D140782484", "courier": "DTDC Express", "amount": 85, "status": "DISPATCHED"},
        {"bib": "RAR-310881", "awb": "7D140782486", "courier": "DTDC Express", "amount": 75, "status": "DISPATCHED"},
        {"bib": "RAR-226585", "awb": "7D140782489", "courier": "DTDC Express", "amount": 98, "status": "DISPATCHED"},
        {"bib": "RAR-474191", "awb": "—",           "courier": "—",            "amount": 78, "status": "PENDING (Blank AWB)"},
        {"bib": "RAR-937297", "awb": "7D140782493", "courier": "DTDC Express", "amount": 75, "status": "DISPATCHED"},
        {"bib": "RAR-677241", "awb": "UNSERVICEABLE","courier": "—",            "amount": "NO", "status": "PENDING (Unserviceable Pin)"},
        {"bib": "RAR-164143", "awb": "7D140782508", "courier": "DTDC Express", "amount": 68, "status": "DISPATCHED"},
        {"bib": "RAR-284014", "awb": "7D140782510", "courier": "DTDC Express", "amount": 100, "status": "DISPATCHED"},
        {"bib": "RAR-437131", "awb": "7D140782511", "courier": "DTDC Express", "amount": 85, "status": "DISPATCHED"},
        {"bib": "RAR-346951", "awb": "UNSERVICEABLE","courier": "—",            "amount": "NO", "status": "PENDING (Unserviceable Pin)"},
        {"bib": "RAR-844382", "awb": "7D140782515", "courier": "DTDC Express", "amount": 100, "status": "DISPATCHED"},
        {"bib": "RAR-326920", "awb": "7D140782516", "courier": "DTDC Express", "amount": 100, "status": "DISPATCHED"},
        {"bib": "RAR-525812", "awb": "7D140782519", "courier": "DTDC Express", "amount": 100, "status": "DISPATCHED"}
    ]
    saini4_map = {x["bib"].upper(): x for x in saini4_data}

    # 4. Fetch full runner details from DB
    conn = mm.get_db_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)

    all_bibs = list(b1_map.keys()) + list(b2_map.keys()) + list(saini4_map.keys())
    # remove duplicates if any
    unique_bibs = []
    seen = set()
    for b in all_bibs:
        if b not in seen:
            unique_bibs.append(b)
            seen.add(b)

    cur.execute("""
        SELECT 
            r.id as registration_id,
            r."bibNumber" as bib_number,
            r."userId" as user_id,
            u.name as user_name,
            u.email as email,
            u.phone as user_phone,
            r."shippingName" as shipping_name,
            r."shippingPhone" as shipping_phone,
            r."shippingLine1" as address_line1,
            r."shippingLine2" as address_line2,
            r."shippingCity" as city,
            r."shippingState" as state,
            r."shippingPincode" as pincode,
            r.distance as distance,
            r."activityType" as activity_type,
            r."finishTimeSeconds" as finish_time_seconds,
            r."proofStatus" as proof_status,
            r."registeredAt" as registered_at,
            m.status as medal_status,
            m."trackingNumber" as tracking_number,
            m.courier as courier,
            m."dispatchedAt" as dispatched_at,
            c."certificateNumber" as certificate_number,
            pu."activityImageUrl" as proof_url
        FROM "Registration" r
        JOIN "User" u ON r."userId" = u.id
        LEFT JOIN "MedalDelivery" m ON m."registrationId" = r.id
        LEFT JOIN "Certificate" c ON c."registrationId" = r.id
        LEFT JOIN "ProofUpload" pu ON pu."registrationId" = r.id
        WHERE r."bibNumber" = ANY(%s)
        ORDER BY r."registeredAt" ASC
    """, (unique_bibs,))
    db_rows = cur.fetchall()
    conn.close()

    db_map = {r["bib_number"].strip().upper(): r for r in db_rows}
    print(f"Fetched {len(db_rows)} athlete records from Database.")

    # Prepare complete records
    saini2_records = []
    saini4_records = []
    combined_records = []

    # Process SAINI 2 (Batch 1 & Batch 2)
    saini2_bib_list = list(b1_map.keys()) + [b for b in b2_map.keys() if b not in b1_map]
    for b in saini2_bib_list:
        db_rec = db_map.get(b)
        b_info = b2_map.get(b) or b1_map.get(b)
        batch_tag = "SAINI 2 (Batch 1)" if b in b1_map else "SAINI 2 (Batch 2)"

        awb = b_info.get("tracking") or (db_rec["tracking_number"] if db_rec else "") or ""
        courier = b_info.get("courier") or (db_rec["courier"] if db_rec else "") or "DTDC Express"
        status = "DISPATCHED" if awb else "PENDING"
        disp_date = "21-Sep-2026"

        rec = {
            "batch": batch_tag,
            "bib": b,
            "db_rec": db_rec,
            "awb": awb,
            "courier": courier,
            "amount": "",
            "status": status,
            "dispatch_date": disp_date
        }
        saini2_records.append(rec)
        combined_records.append(rec)

    # Process SAINI 4 (Batch 3)
    for s in saini4_data:
        b = s["bib"]
        db_rec = db_map.get(b)
        batch_tag = "SAINI 4 (Batch 3)"
        disp_date = "22-Sep-2026" if s["status"] == "DISPATCHED" else "Pending"

        rec = {
            "batch": batch_tag,
            "bib": b,
            "db_rec": db_rec,
            "awb": s["awb"],
            "courier": s["courier"],
            "amount": s["amount"],
            "status": s["status"],
            "dispatch_date": disp_date
        }
        saini4_records.append(rec)
        combined_records.append(rec)

    print(f"Total Combined Records: {len(combined_records)} (SAINI 2: {len(saini2_records)}, SAINI 4: {len(saini4_records)})")

    # 5. Build Workbook
    wb = openpyxl.Workbook()

    # Styling constants
    DARK_GREEN = "0F3D24"
    MED_GREEN  = "1B633D"
    SUB_HEADER = "2D7D52"
    ZEBRA_BG   = "F7FAF8"
    WHITE_BG   = "FFFFFF"
    BORDER_CLR = "D1E2D8"

    fill_banner = PatternFill(start_color=DARK_GREEN, end_color=DARK_GREEN, fill_type="solid")
    fill_header = PatternFill(start_color=MED_GREEN, end_color=MED_GREEN, fill_type="solid")
    fill_zebra  = PatternFill(start_color=ZEBRA_BG, end_color=ZEBRA_BG, fill_type="solid")
    fill_white  = PatternFill(start_color=WHITE_BG, end_color=WHITE_BG, fill_type="solid")

    fill_disp = PatternFill(start_color="DEF7EC", end_color="DEF7EC", fill_type="solid")  # Green pill
    font_disp = Font(name="Segoe UI", size=9, bold=True, color="03543F")

    fill_pend = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # Amber pill
    font_pend = Font(name="Segoe UI", size=9, bold=True, color="92400E")

    fill_batch_s2 = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")  # Light blue
    font_batch_s2 = Font(name="Segoe UI", size=9, bold=True, color="0369A1")

    fill_batch_s4 = PatternFill(start_color="FCE7F3", end_color="FCE7F3", fill_type="solid")  # Light pink
    font_batch_s4 = Font(name="Segoe UI", size=9, bold=True, color="BE185D")

    thin_border = Border(
        left=Side(style="thin", color=BORDER_CLR),
        right=Side(style="thin", color=BORDER_CLR),
        top=Side(style="thin", color=BORDER_CLR),
        bottom=Side(style="thin", color=BORDER_CLR)
    )

    f_banner_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    f_banner_sub   = Font(name="Segoe UI", size=9.5, italic=True, color="E2E8F0")
    f_header       = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
    f_bold         = Font(name="Segoe UI", size=9.5, bold=True, color="111827")
    f_regular      = Font(name="Segoe UI", size=9.5, color="111827")
    f_link         = Font(name="Segoe UI", size=9.5, color="1D4ED8", underline="single")

    align_center = Alignment(horizontal="center", vertical="center")
    align_left   = Alignment(horizontal="left", vertical="center")
    align_wrap   = Alignment(horizontal="left", vertical="center", wrap_text=True)

    COLS = [
        ("S.No", 6, align_center),
        ("Batch Source", 16, align_center),
        ("BIB Number", 13, align_center),
        ("Recipient Name", 22, align_left),
        ("Primary Mobile", 15, align_center),
        ("Email Address", 26, align_left),
        ("Full Shipping Address", 44, align_wrap),
        ("City", 16, align_left),
        ("State", 18, align_left),
        ("Pincode", 10, align_center),
        ("Challenge Distance", 18, align_center),
        ("Package Contents", 26, align_left),
        ("Courier Partner", 16, align_center),
        ("AWB / Tracking No", 18, align_center),
        ("Tracking Link", 30, align_left),
        ("Courier Amount (₹)", 15, align_center),
        ("Dispatch Date", 14, align_center),
        ("Delivery Status", 18, align_center),
    ]

    def render_sheet(ws, title, subtitle, records):
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "E4"

        last_col = get_column_letter(len(COLS))

        # Title
        ws.merge_cells(f"A1:{last_col}1")
        ws["A1"] = title
        ws["A1"].font = f_banner_title
        ws["A1"].fill = fill_banner
        ws["A1"].alignment = align_center
        ws.row_dimensions[1].height = 30

        # Subtitle
        ws.merge_cells(f"A2:{last_col}2")
        ws["A2"] = subtitle
        ws["A2"].font = f_banner_sub
        ws["A2"].fill = PatternFill(start_color=MED_GREEN, end_color=MED_GREEN, fill_type="solid")
        ws["A2"].alignment = align_center
        ws.row_dimensions[2].height = 20

        # Header
        ws.row_dimensions[3].height = 25
        for col_idx, (col_name, col_width, col_align) in enumerate(COLS, 1):
            letter = get_column_letter(col_idx)
            ws.column_dimensions[letter].width = col_width
            cell = ws.cell(row=3, column=col_idx, value=col_name)
            cell.font = f_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border

        # Data rows
        for idx, item in enumerate(records, 1):
            r_idx = idx + 3
            ws.row_dimensions[r_idx].height = 22
            is_even = (idx % 2 == 0)
            row_fill = fill_zebra if is_even else fill_white

            db = item["db_rec"] or {}
            rec_name = db.get("shipping_name") or db.get("user_name") or "Runner"
            phone = db.get("shipping_phone") or db.get("user_phone") or ""
            email = db.get("email") or ""
            city = (db.get("city") or "").title()
            state = (db.get("state") or "").title()
            pincode = db.get("pincode") or ""

            line1 = db.get("address_line1") or ""
            line2 = db.get("address_line2") or ""
            addr_parts = [p for p in [line1, line2] if p]
            full_addr = f"{', '.join(addr_parts)}, {city}, {state} - {pincode}".strip(" ,-")

            dist = db.get("distance") or "5 KM"
            dist_str = f"{dist} Challenge" if not str(dist).endswith("Challenge") else dist

            awb = item["awb"]
            courier = item["courier"] or "DTDC Express"
            track_url = ""
            if awb and awb not in ["—", "UNSERVICEABLE"]:
                if "428" in awb:
                    track_url = f"https://www.delhivery.com/track/package/{awb}"
                else:
                    track_url = "https://www.dtdc.com/track-your-shipment/"

            values = [
                idx,
                item["batch"],
                item["bib"],
                rec_name,
                phone,
                email,
                full_addr,
                city,
                state,
                pincode,
                dist_str,
                "Rise & Run Finisher Medal",
                courier,
                awb,
                track_url,
                item["amount"] if item["amount"] else "—",
                item["dispatch_date"],
                item["status"]
            ]

            for c_idx, val in enumerate(values, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = val
                cell.border = thin_border
                cell.fill = row_fill
                align_rule = COLS[c_idx - 1][2]
                cell.alignment = align_rule
                cell.font = f_regular

                # Specific cell formatting
                if c_idx == 2:  # Batch Source
                    if "SAINI 2" in str(val):
                        cell.fill = fill_batch_s2
                        cell.font = font_batch_s2
                    else:
                        cell.fill = fill_batch_s4
                        font_batch_s4
                elif c_idx == 3:  # BIB
                    cell.font = f_bold
                elif c_idx == 4:  # Name
                    cell.font = f_bold
                elif c_idx == 10: # Pincode
                    cell.font = f_bold
                elif c_idx == 14: # AWB
                    cell.font = f_bold
                elif c_idx == 15 and val: # Tracking Link
                    cell.hyperlink = val
                    cell.font = f_link
                elif c_idx == 18: # Delivery Status
                    if "DISPATCHED" in str(val):
                        cell.fill = fill_disp
                        cell.font = font_disp
                    else:
                        cell.fill = fill_pend
                        cell.font = font_pend

    # --- Sheet 1: Master Combined (All 134) ---
    ws1 = wb.active
    ws1.title = "Master Combined List"
    render_sheet(
        ws1,
        "MOUNTAIN RUN  |  RISE AND RUN VIRTUAL CHALLENGE 2026",
        f"Master Consolidate Dispatch List (SAINI 2 + SAINI 4)  |  Total: {len(combined_records)} Athletes  |  Generated: {datetime.datetime.now().strftime('%d-%b-%Y %I:%M %p')}",
        combined_records
    )

    # --- Sheet 2: SAINI 2 ---
    ws2 = wb.create_sheet(title="SAINI 2 (Batch 1 & 2)")
    render_sheet(
        ws2,
        "MOUNTAIN RUN  |  SAINI 2 DISPATCH BATCH",
        f"SAINI 2 Athlete Dispatch List  |  Total: {len(saini2_records)} Athletes Dispatched",
        saini2_records
    )

    # --- Sheet 3: SAINI 4 ---
    ws3 = wb.create_sheet(title="SAINI 4 (Batch 3)")
    render_sheet(
        ws3,
        "MOUNTAIN RUN  |  SAINI 4 DISPATCH BATCH",
        f"SAINI 4 Athlete Dispatch List  |  Total: {len(saini4_records)} Athletes (16 Dispatched, 3 Pending)",
        saini4_records
    )

    # --- Sheet 4: Summary KPI ---
    ws4 = wb.create_sheet(title="Summary & Analytics")
    ws4.views.sheetView[0].showGridLines = True
    ws4.column_dimensions["B"].width = 28
    ws4.column_dimensions["C"].width = 24
    ws4.column_dimensions["D"].width = 32

    ws4.merge_cells("B2:D2")
    ws4["B2"] = "RISE AND RUN — DISPATCH CONSOLIDATION KPI"
    ws4["B2"].font = f_banner_title
    ws4["B2"].fill = fill_banner
    ws4["B2"].alignment = align_center
    ws4.row_dimensions[2].height = 28

    kpis = [
        ("Total Unique Athletes", len(combined_records)),
        ("Total Successfully Dispatched", sum(1 for r in combined_records if "DISPATCHED" in r["status"])),
        ("Total Pending / Unserviceable", sum(1 for r in combined_records if "PENDING" in r["status"])),
        ("SAINI 2 Batch Dispatched", len(saini2_records)),
        ("SAINI 4 Batch Dispatched", sum(1 for r in saini4_records if "DISPATCHED" in r["status"])),
        ("SAINI 4 Batch Pending", sum(1 for r in saini4_records if "PENDING" in r["status"])),
        ("Primary Courier Partner", "DTDC Express (98.5%)"),
        ("Secondary Courier Partner", "Delhivery Express (1.5%)"),
        ("Email Dispatch Status", "100% of Dispatched Athletes have received Tracking Emails via Resend")
    ]

    for i, (k, v) in enumerate(kpis, 4):
        ws4.cell(row=i, column=2, value=k).font = f_bold
        ws4.cell(row=i, column=2).border = thin_border
        ws4.cell(row=i, column=2).fill = fill_zebra
        ws4.cell(row=i, column=3, value=v).font = f_regular
        ws4.cell(row=i, column=3).alignment = align_center
        ws4.cell(row=i, column=3).border = thin_border
        ws4.cell(row=i, column=3).fill = fill_white
        ws4.row_dimensions[i].height = 22

    # Save to root as SAINI_MASTER_COMBINED.xlsx and SAINI_COMBINED.xlsx
    out_name = "SAINI_MASTER_COMBINED.xlsx"
    out_path = os.path.join(os.getcwd(), out_name)
    wb.save(out_path)

    copy_name = "SAINI_COMBINED.xlsx"
    copy_path = os.path.join(os.getcwd(), copy_name)
    wb.save(copy_path)

    print(f"\n✅ Successfully generated:")
    print(f"  1. {out_path}")
    print(f"  2. {copy_path}")

    return out_path

if __name__ == "__main__":
    build_combined_saini_excel()
